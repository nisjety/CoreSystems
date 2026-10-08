# Reference program for the H5 "docx" case. See programs/make_xlsx.py.
from docx import Document

doc = Document()
doc.add_heading("Statusnotat", level=0)
doc.add_paragraph("Prosjektet ligger an til å levere i tide. Alle hovedaktivitetene er i gang.")
doc.add_paragraph("Neste uke starter testingen, og vi følger opp risikoene på det ukentlige møtet.")

table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
for cell, text in zip(table.rows[0].cells, ["Aktivitet", "Ansvarlig", "Frist"]):
    cell.text = text
for activity, owner, due in [
    ("Kravspesifikasjon", "Kari", "12. oktober"),
    ("Prototype", "Ola", "26. oktober"),
    ("Brukertest", "Nora", "9. november"),
]:
    cells = table.add_row().cells
    cells[0].text, cells[1].text, cells[2].text = activity, owner, due

doc.save("statusnotat.docx")
print("saved statusnotat.docx")
