# Reference program for the H5 "pptx" case: a matplotlib PNG placed on a slide with
# python-pptx. See programs/make_xlsx.py.
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches

months = ["Juli", "August", "September"]
revenue = [120, 135, 150]
fig, ax = plt.subplots(figsize=(6, 3.6), dpi=100)
ax.bar(months, revenue, color="#2f6f9f")
ax.set_title("Omsetning per måned")
ax.set_ylabel("Millioner NOK")
fig.tight_layout()
fig.savefig("omsetning.png")
plt.close(fig)

prs = Presentation()
title_slide = prs.slides.add_slide(prs.slide_layouts[0])
title_slide.shapes.title.text = "Kvartalsrapport"
title_slide.placeholders[1].text = "Tredje kvartal"

chart_slide = prs.slides.add_slide(prs.slide_layouts[5])
chart_slide.shapes.title.text = "Omsetning per måned"
chart_slide.shapes.add_picture("omsetning.png", Inches(1.5), Inches(1.8), width=Inches(7))

prs.save("kvartalsrapport.pptx")
print("saved kvartalsrapport.pptx")
