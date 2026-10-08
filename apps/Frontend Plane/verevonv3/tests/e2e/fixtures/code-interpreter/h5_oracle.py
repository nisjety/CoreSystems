#!/usr/bin/env python3
"""Oracle for the H5 live proof: do the files code_interpreter hands back hold what was asked?

The live spec (tests/e2e/product-code-interpreter.spec.ts) saves every file a Terra
turn delivered and asks this script about it. It reads each file with the real
library for its type, the same libraries the sandbox image ships (execution-core's
Dockerfile): openpyxl for .xlsx, python-docx for .docx, pypdf for .pdf,
python-pptx for .pptx and Pillow for .png. A file that merely exists proves nothing:
a workbook with a chart but no formulas, a deck with no picture, a PDF whose
Norwegian letters came out as boxes and a chart that is a blank white square all
exist.

Commands (each prints one JSON document on stdout):

  check <dir>             verdict per file in <dir>, by extension
  reference               run programs/make_*.py in a throwaway directory and print
                          the files they wrote, in code_interpreter's own shape
                          ({name, mime, bytes, content_b64}) with per-program timings
  self-test               reference, then check, then break each file in a known way
                          and require the oracle to notice. Exit 1 if it does not.

Run it where the libraries are, e.g. in the stack's own image (see the runbook):

  docker run --rm --network none -v "<this directory>:/h5:ro" \\
      --entrypoint python3 model-plane-execution-core /h5/h5_oracle.py self-test

Only the standard library is imported at the top: the libraries load lazily, so a
missing one is reported per file instead of crashing the whole run.
"""
import base64
import json
import os
import re
import subprocess
import sys
import tempfile
import time
import unicodedata
import zipfile

HERE = os.path.dirname(os.path.abspath(globals().get("__file__", ".")))  # no __file__ when piped to `python3 -`
PROGRAMS = os.path.join(HERE, "programs")

MIME = {
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".pdf": "application/pdf",
    ".png": "image/png",
}
MAGIC = {
    ".xlsx": b"PK\x03\x04",
    ".docx": b"PK\x03\x04",
    ".pptx": b"PK\x03\x04",
    ".pdf": b"%PDF-",
    ".png": b"\x89PNG\r\n\x1a\n",
}

# What the five prompts in tests/e2e/code-interpreter-proof.ts ask for. Words the
# prompts spell out are checked as spelled; anything the model is free to choose is
# only checked structurally.
PDF_WORDS = ["Sammendrag", "blåbær", "grøt", "Ålesund"]
DOCX_HEADER = ["aktivitet", "ansvarlig", "frist"]
XLSX_HEADER = ["omsetning", "kostnad", "margin"]
CHART_KINDS = "bar|line|pie|area|scatter|doughnut|radar|bubble|stock|surface"


def nfc(text):
    return unicodedata.normalize("NFC", text or "")


def folded(text):
    return nfc(text).casefold()


# ---------------------------------------------------------------------------
# Pixels, shared by the .png case and the picture on a slide.
# ---------------------------------------------------------------------------


def pixel_facts(image):
    """Distinct colours and how much of the picture is its dominant colour.

    A transparent picture is flattened on white first, the way a viewer shows it,
    so a chart saved with a transparent background is not read as a black square.
    """
    from PIL import Image

    if image.mode in ("RGBA", "LA") or (image.mode == "P" and "transparency" in image.info):
        rgba = image.convert("RGBA")
        flat = Image.new("RGB", rgba.size, "white")
        flat.paste(rgba, mask=rgba.split()[-1])
    else:
        flat = image.convert("RGB")
    colors = flat.getcolors(maxcolors=flat.width * flat.height) or []
    total = sum(count for count, _ in colors) or 1
    dominant = max((count for count, _ in colors), default=0)
    return {
        "width": flat.width,
        "height": flat.height,
        "distinct_colors": len(colors),
        "dominant_fraction": round(dominant / total, 4),
    }


def pixel_problems(facts, minimum_width, minimum_height):
    problems = []
    if facts["width"] < minimum_width or facts["height"] < minimum_height:
        problems.append("picture is only %dx%d" % (facts["width"], facts["height"]))
    # A chart on a white page is mostly white, so the bar is high; a blank picture
    # is one colour.
    if facts["distinct_colors"] < 8 or facts["dominant_fraction"] > 0.995:
        problems.append(
            "picture looks blank (%d colours, %.1f%% one colour)"
            % (facts["distinct_colors"], facts["dominant_fraction"] * 100)
        )
    return problems


# ---------------------------------------------------------------------------
# One checker per type. Each returns (facts, problems).
# ---------------------------------------------------------------------------


def check_xlsx(path):
    import openpyxl

    facts, problems = {}, []
    with zipfile.ZipFile(path) as archive:
        chart_parts = sorted(n for n in archive.namelist() if re.fullmatch(r"xl/charts/chart\d*\.xml", n))
        facts["chart_parts"] = len(chart_parts)
        kinds, references = set(), 0
        for name in chart_parts:
            xml = archive.read(name).decode("utf-8", "replace")
            kinds.update(re.findall(r"<(?:\w+:)?(%s)Chart\b" % CHART_KINDS, xml))
            references += len(re.findall(r"<(?:\w+:)?f>[^<]*!", xml))
        facts["chart_kinds"] = sorted(kinds)
        facts["chart_cell_references"] = references
        facts["drawing_parts"] = sum(1 for n in archive.namelist() if re.fullmatch(r"xl/drawings/drawing\d*\.xml", n))
    if not chart_parts:
        problems.append("no native chart part (xl/charts/chart*.xml)")
    elif "bar" not in kinds:
        problems.append("the chart is not a bar chart (found: %s)" % (", ".join(sorted(kinds)) or "unknown"))
    if chart_parts and references == 0:
        problems.append("the chart is not bound to any cell range")
    if chart_parts and facts["drawing_parts"] == 0:
        problems.append("a chart part exists but no drawing places it on a sheet")

    workbook = openpyxl.load_workbook(path)  # formulas as written, not their cached values
    facts["sheets"] = workbook.sheetnames
    formulas = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    formulas.append(str(cell.value))
    facts["formula_cells"] = len(formulas)
    facts["sum_formulas"] = sum(1 for f in formulas if f.upper().startswith("=SUM("))
    if len(formulas) < 4:
        problems.append("%d formula cells; 3 per-row margins and a sum row were asked for" % len(formulas))
    if not facts["sum_formulas"]:
        problems.append("no SUM formula")
    first = workbook.worksheets[0]
    header = [folded(str(c.value)) for c in next(first.iter_rows(min_row=1, max_row=1), []) if c.value is not None]
    facts["header_columns"] = len(header)
    for word in XLSX_HEADER:
        if not any(word in column for column in header):
            problems.append("header row has no '%s' column" % word)
    try:
        facts["charts_loaded_by_openpyxl"] = sum(len(sheet._charts) for sheet in workbook.worksheets)
    except Exception:  # private attribute; the zip check above is the authority
        facts["charts_loaded_by_openpyxl"] = None
    return facts, problems


def check_docx(path):
    from docx import Document

    document = Document(path)
    facts, problems = {}, []
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    facts["paragraphs"] = len(paragraphs)
    facts["tables"] = len(document.tables)
    if facts["tables"] < 1:
        problems.append("no table")
    else:
        table = document.tables[0]
        facts["table_rows"] = len(table.rows)
        facts["table_columns"] = len(table.columns)
        header = [folded(cell.text) for cell in table.rows[0].cells]
        for word in DOCX_HEADER:
            if not any(word in column for column in header):
                problems.append("table header has no '%s' column" % word)
        if len(table.rows) < 4:
            problems.append("table has %d rows; a header and three rows were asked for" % len(table.rows))
        if len(table.columns) < 3:
            problems.append("table has %d columns; three were asked for" % len(table.columns))
    if len(paragraphs) < 3:
        problems.append("%d paragraphs outside the table; a title and two paragraphs were asked for" % len(paragraphs))
    if not any("statusnotat" in folded(p) for p in paragraphs):
        problems.append("the title 'Statusnotat' is not in the document")
    return facts, problems


def check_pdf(path):
    from pypdf import PdfReader

    reader = PdfReader(path)
    facts, problems = {}, []
    facts["pages"] = len(reader.pages)
    text = nfc("\n".join((page.extract_text() or "") for page in reader.pages))
    facts["text_chars"] = len(text)
    facts["has_replacement_character"] = "�" in text
    if facts["pages"] < 1:
        problems.append("no pages")
    for word in PDF_WORDS:
        if nfc(word) not in text:
            problems.append("the word '%s' is not in the extracted text" % word)
    for letter in "æøå":
        facts["has_" + letter] = letter in text
    if facts["has_replacement_character"]:
        problems.append("the text holds a replacement character")
    return facts, problems


def check_pptx(path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from PIL import Image
    import io

    presentation = Presentation(path)
    facts, problems = {}, []
    slides = list(presentation.slides)
    facts["slides"] = len(slides)
    pictures, texts = [], []

    def walk(shapes):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(shape.shapes)
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures.append(shape)
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                texts.append(shape.text_frame.text)

    slides_with_picture = 0
    for slide in slides:
        before = len(pictures)
        walk(slide.shapes)
        if len(pictures) > before:
            slides_with_picture += 1
    facts["pictures"] = len(pictures)
    facts["slides_with_picture"] = slides_with_picture
    if len(slides) < 2:
        problems.append("%d slides; two were asked for" % len(slides))
    if not pictures:
        problems.append("no picture on any slide")
    else:
        stats = pixel_facts(Image.open(io.BytesIO(pictures[0].image.blob)))
        facts["first_picture"] = stats
        problems.extend("the picture on the slide: " + p for p in pixel_problems(stats, 100, 100))
    if not any("kvartalsrapport" in folded(t) for t in texts):
        problems.append("the title 'Kvartalsrapport' is not on a slide")
    return facts, problems


def check_png(path):
    from PIL import Image

    with Image.open(path) as image:
        facts = {"format": image.format}
        image.load()
        facts.update(pixel_facts(image))
    problems = []
    if facts["format"] != "PNG":
        problems.append("the file is %s, not PNG" % facts["format"])
    problems.extend(pixel_problems(facts, 300, 200))
    return facts, problems


CHECKERS = {
    ".xlsx": check_xlsx,
    ".docx": check_docx,
    ".pdf": check_pdf,
    ".pptx": check_pptx,
    ".png": check_png,
}


def check_file(path):
    extension = os.path.splitext(path)[1].lower()
    result = {"file": os.path.basename(path), "kind": extension.lstrip("."), "ok": False, "problems": [], "facts": {}}
    try:
        with open(path, "rb") as handle:
            head = handle.read(16)
        size = os.path.getsize(path)
        result["bytes"] = size
        if size == 0:
            result["problems"].append("the file is empty")
            return result
        if not head.startswith(MAGIC[extension]):
            result["problems"].append("the bytes are not a %s file" % extension)
            return result
        facts, problems = CHECKERS[extension](path)
        result["facts"], result["problems"] = facts, problems
        result["ok"] = not problems
    except Exception as error:  # a file the library cannot open is a verdict, not a crash
        result["problems"].append("the library could not read the file: %s" % type(error).__name__)
    return result


def check_directory(directory):
    results = []
    for name in sorted(os.listdir(directory)):
        if os.path.splitext(name)[1].lower() in CHECKERS:
            results.append(check_file(os.path.join(directory, name)))
    return {"results": results}


# ---------------------------------------------------------------------------
# reference: the five programs, run the way the sandbox runs a call.
# ---------------------------------------------------------------------------


def run_reference(destination=None):
    """Run programs/make_*.py, each in its own empty directory (as the sandbox does).

    With `destination`, the files are kept there; otherwise the caller gets them
    base64-encoded and the directories are removed.
    """
    produced, timings = [], []
    environment = dict(os.environ, MPLBACKEND="Agg", PYTHONDONTWRITEBYTECODE="1")
    for program in sorted(n for n in os.listdir(PROGRAMS) if n.startswith("make_") and n.endswith(".py")):
        with tempfile.TemporaryDirectory(prefix="h5-ref-") as workdir:
            environment["MPLCONFIGDIR"] = os.path.join(workdir, ".mplconfig")
            os.makedirs(environment["MPLCONFIGDIR"])
            code = open(os.path.join(PROGRAMS, program), "rb").read()
            staged = os.path.join(workdir, "_verevon_main.py")  # the sandbox's own program name
            with open(staged, "wb") as handle:
                handle.write(code)
            started = time.time()
            run = subprocess.run(
                [sys.executable, "_verevon_main.py"], cwd=workdir, env=environment,
                capture_output=True, timeout=60,
            )
            elapsed = round(time.time() - started, 2)
            timings.append({"program": program, "exit_code": run.returncode, "seconds": elapsed,
                            "stderr_tail": run.stderr.decode("utf-8", "replace")[-300:]})
            for name in sorted(os.listdir(workdir)):
                full = os.path.join(workdir, name)
                if name.startswith(".") or name == "_verevon_main.py" or not os.path.isfile(full):
                    continue
                data = open(full, "rb").read()
                extension = os.path.splitext(name)[1].lower()
                produced.append({
                    "name": name,
                    "mime": MIME.get(extension, "application/octet-stream"),
                    "bytes": len(data),
                    "content_b64": base64.b64encode(data).decode("ascii"),
                    "program": program,
                })
                if destination:
                    os.makedirs(destination, exist_ok=True)
                    with open(os.path.join(destination, name), "wb") as handle:
                        handle.write(data)
    return {"files": produced, "timings": timings}


# ---------------------------------------------------------------------------
# self-test: the oracle must accept the reference files AND reject broken ones.
# ---------------------------------------------------------------------------


def broken_variants(good_dir, scratch):
    """(label, expected problem substring, path) for files broken in one known way."""
    import openpyxl
    from openpyxl.chart import BarChart, Reference
    from docx import Document
    from PIL import Image
    from pptx import Presentation
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    def path(name):
        return os.path.join(scratch, name)

    variants = []

    # A workbook with formulas but no chart, and one with a chart but no formulas.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Måned", "Omsetning", "Kostnad", "Margin"])
    for index, (month, revenue, cost) in enumerate([("Juli", 1, 1), ("August", 2, 1), ("September", 3, 1)], start=2):
        ws.append([month, revenue, cost, "=B%d-C%d" % (index, index)])
    ws.append(["Sum", "=SUM(B2:B4)", "=SUM(C2:C4)", "=SUM(D2:D4)"])
    wb.save(path("no_chart.xlsx"))
    variants.append(("xlsx without a chart", "no native chart", path("no_chart.xlsx")))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Måned", "Omsetning", "Kostnad", "Margin"])
    for month, revenue, cost in [("Juli", 1, 1), ("August", 2, 1), ("September", 3, 1)]:
        ws.append([month, revenue, cost, revenue - cost])
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    ws.add_chart(chart, "F2")
    wb.save(path("values_only.xlsx"))
    variants.append(("xlsx with a chart but only typed numbers", "formula cells", path("values_only.xlsx")))

    # A document without a table.
    document = Document()
    document.add_heading("Statusnotat", level=0)
    document.add_paragraph("Ett avsnitt.")
    document.add_paragraph("To avsnitt.")
    document.save(path("no_table.docx"))
    variants.append(("docx without a table", "no table", path("no_table.docx")))

    # A PDF whose Norwegian letters were replaced by ASCII.
    pdf = canvas.Canvas(path("ascii.pdf"), pagesize=A4)
    pdf.drawString(72, 700, "Sammendrag")
    pdf.drawString(72, 680, "Rod grot med flote er sot og sunn, og blabaer fra Alesund smaker best.")
    pdf.save()
    variants.append(("pdf without the Norwegian letters", "is not in the extracted text", path("ascii.pdf")))

    # A deck with two slides and no picture.
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[0])
    first.shapes.title.text = "Kvartalsrapport"
    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    second.shapes.title.text = "Omsetning per måned"
    presentation.save(path("no_picture.pptx"))
    variants.append(("pptx without a picture", "no picture on any slide", path("no_picture.pptx")))

    # A blank white chart.
    Image.new("RGB", (640, 480), "white").save(path("blank.png"))
    variants.append(("a blank white png", "looks blank", path("blank.png")))

    # A file cut short, and an empty one.
    good = open(os.path.join(good_dir, "salg-q3.xlsx"), "rb").read()
    with open(path("cut.xlsx"), "wb") as handle:
        handle.write(good[: len(good) // 3])
    variants.append(("an xlsx cut off after a third", "could not read", path("cut.xlsx")))
    open(path("empty.pdf"), "wb").close()
    variants.append(("an empty pdf", "is empty", path("empty.pdf")))
    return variants


def self_test():
    cases, failures = [], []
    with tempfile.TemporaryDirectory(prefix="h5-self-") as root:
        good_dir = os.path.join(root, "good")
        reference = run_reference(good_dir)
        bad_runs = [t for t in reference["timings"] if t["exit_code"] != 0]
        for run in bad_runs:
            failures.append("%s exited %d: %s" % (run["program"], run["exit_code"], run["stderr_tail"]))
        verdicts = check_directory(good_dir)["results"]
        expected = {"salg-q3.xlsx", "statusnotat.docx", "sammendrag.pdf", "kvartalsrapport.pptx", "besokstall.png"}
        seen = {v["file"] for v in verdicts}
        for missing in sorted(expected - seen):
            failures.append("the reference produced no %s" % missing)
        for verdict in verdicts:
            cases.append({"case": "accepts " + verdict["file"], "ok": verdict["ok"], "problems": verdict["problems"]})
            if not verdict["ok"]:
                failures.append("rejected a reference file %s: %s" % (verdict["file"], verdict["problems"]))
        scratch = os.path.join(root, "broken")
        os.makedirs(scratch)
        for label, needle, broken in broken_variants(good_dir, scratch):
            verdict = check_file(broken)
            noticed = (not verdict["ok"]) and any(needle in p for p in verdict["problems"])
            cases.append({"case": "rejects " + label, "ok": noticed, "problems": verdict["problems"]})
            if not noticed:
                failures.append("did not notice: %s (wanted '%s', got %s)" % (label, needle, verdict["problems"]))
    return {"ok": not failures, "failures": failures, "cases": cases, "timings": reference["timings"]}


def main(argv):
    command = argv[1] if len(argv) > 1 else ""
    if command == "check" and len(argv) == 3:
        document = check_directory(argv[2])
    elif command == "reference" and len(argv) == 2:
        document = run_reference()
    elif command == "self-test" and len(argv) == 2:
        document = self_test()
    else:
        sys.stderr.write(__doc__)
        return 2
    sys.stdout.write(json.dumps(document))
    sys.stdout.write("\n")
    return 0 if document.get("ok", True) else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv))
