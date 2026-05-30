"""
Documents API — v1 routes.

POST   /v1/documents              Ingest a new document
GET    /v1/documents              List documents for an org
GET    /v1/documents/{id}         Get a single document (with content)
DELETE /v1/documents/{id}         Delete document (triggers vector cleanup)
PATCH  /v1/documents/{id}/status  Internal: update processing status
"""
from __future__ import annotations

import hmac
import logging
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Header, Query, Request, status
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.postgres import (
    create_document,
    delete_document,
    get_document,
    list_documents,
    search_documents,
    update_document_status,
    get_db,
    get_document_content,
)
from app.events.publisher import publish_document_created, publish_document_deleted
from app.quota_enforcement_handler import get_quota_blocker
from app.models.document import (
    DocumentCreateRequest,
    DocumentListResponse,
    DocumentResponse,
    DocumentSearchItem,
    DocumentSearchResponse,
)

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/v1/documents", tags=["documents"])


# ── Auth helpers ──────────────────────────────────────────────────────────────

def _get_auth(request: Request):
    """Extract AuthContext from request.state (set by auth middleware)."""
    from shared.auth_middleware import get_auth_context
    return get_auth_context(request)


def _get_org_access(request: Request):
    """Extract OrgAccess from request.state (set by authz middleware)."""
    access = getattr(request.state, "org_access", None)
    if access is None:
        raise HTTPException(status_code=403, detail="Organization access not resolved")
    return access


# ── Internal auth for worker-to-service calls ─────────────────────────────────

def _verify_internal_key(x_service_auth: Optional[str] = Header(None)):
    """Validate X-Service-Auth header for internal (worker → service) calls."""
    from app.config import settings
    if not settings.internal_api_key:
        raise HTTPException(status_code=503, detail="Internal auth not configured")
    if not hmac.compare_digest(x_service_auth or "", settings.internal_api_key):
        raise HTTPException(status_code=403, detail="Invalid service auth key")
    return True


# ── Ingest ────────────────────────────────────────────────────────────────────

@router.post("", response_model=DocumentResponse, status_code=status.HTTP_201_CREATED)
async def ingest_document(
    request: Request,
    body: DocumentCreateRequest,
    db: AsyncSession = Depends(get_db),
) -> DocumentResponse:
    """
    Ingest a raw document.
    Saves to Postgres with status=pending, then emits dataplane.documents.created.
    org_id derived from authenticated session — NOT from request body.
    """
    auth = _get_auth(request)
    access = _get_org_access(request)

    # Permission check: requires resources:create
    if not access.has_permission("resources:create"):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=f"Permission 'resources:create' required. Your role '{access.role}' does not have this.",
        )

    org_id = auth.org_id

    # Quota enforcement: block ingestion when org has exceeded billing limits
    if get_quota_blocker().is_blocked(org_id):
        raise HTTPException(
            status_code=status.HTTP_429_TOO_MANY_REQUESTS,
            detail=f"Organisation {org_id!r} has exceeded its quota limit. "
                   "Upgrade your plan to continue.",
        )

    doc = await create_document(
        db,
        org_id=org_id,
        source=body.source,
        type=body.type,
        title=body.title,
        content=body.content,
        metadata=body.metadata,
        created_by=auth.user_id,
    )
    await publish_document_created(
        document_id=doc["document_id"],
        org_id=doc["org_id"],
        payload={
            "document_id": doc["document_id"],
            "org_id": doc["org_id"],
            "source": doc["source"],
            "type": doc["type"],
            "title": doc["title"],
            "user_id": auth.user_id,
        },
    )
    return DocumentResponse(**doc)


# ── List ──────────────────────────────────────────────────────────────────────

@router.get("", response_model=DocumentListResponse)
async def list_docs(
    request: Request,
    q: Optional[str] = Query(None, description="Filter by title (case-insensitive)"),
    type: Optional[str] = Query(None),
    status_filter: Optional[str] = Query(None, alias="status"),
    limit: int = Query(50, ge=1, le=200),
    offset: int = Query(0, ge=0),
    db: AsyncSession = Depends(get_db),
) -> DocumentListResponse:
    auth = _get_auth(request)
    access = _get_org_access(request)

    if not access.has_permission("resources:read"):
        raise HTTPException(status_code=403, detail="Permission 'resources:read' required")

    docs = await list_documents(
        db, org_id=auth.org_id, q=q, type=type, status=status_filter, limit=limit, offset=offset
    )
    return DocumentListResponse(
        documents=[DocumentResponse(**d) for d in docs],
        total=len(docs),
    )


# ── Search (autocomplete) ──────────────────────────────────────────────────────

@router.get("/search", response_model=DocumentSearchResponse)
async def search_docs(
    request: Request,
    q: str = Query(..., min_length=1, max_length=200, description="Title search query"),
    limit: int = Query(5, ge=1, le=20, description="Max results (default 5, max 20)"),
    db: AsyncSession = Depends(get_db),
) -> DocumentSearchResponse:
    """
    Lightweight document title search for autocomplete and slash-command /file suggestions.
    Only returns completed documents. Excludes content and metadata for minimal payload.
    """
    auth = _get_auth(request)
    access = _get_org_access(request)

    if not access.has_permission("resources:read"):
        raise HTTPException(status_code=403, detail="Permission 'resources:read' required")

    docs = await search_documents(db, org_id=auth.org_id, q=q, limit=limit)
    return DocumentSearchResponse(
        results=[DocumentSearchItem(**d) for d in docs],
        query=q,
        total=len(docs),
    )


# ── Detail ────────────────────────────────────────────────────────────────────

@router.get("/{document_id}", response_model=DocumentResponse)
async def get_doc(
    document_id: str,
    request: Request,
    db: AsyncSession = Depends(get_db),
) -> DocumentResponse:
    auth = _get_auth(request)
    access = _get_org_access(request)

    if not access.has_permission("resources:read"):
        raise HTTPException(status_code=403, detail="Permission 'resources:read' required")

    doc = await get_document(db, document_id=document_id, org_id=auth.org_id)
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")
    return DocumentResponse(**doc)


# ── Export ───────────────────────────────────────────────────────────────────

@router.get("/{document_id}/export")
async def export_document_as_word(
    document_id: str,
    request: Request,
    db: AsyncSession = Depends(get_db),
) -> StreamingResponse:
    """Export a document as a .docx (Word) file."""
    auth = _get_auth(request)
    access = _get_org_access(request)

    if not access.has_permission("resources:read"):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Permission 'resources:read' required",
        )

    doc = await get_document_content(db, document_id=document_id, org_id=auth.org_id)
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")

    from docx import Document as DocxDocument
    from io import BytesIO

    word = DocxDocument()
    word.add_heading(doc.get("title") or "Untitled", level=0)
    if doc.get("content"):
        word.add_paragraph(doc["content"])

    buf = BytesIO()
    word.save(buf)
    buf.seek(0)

    safe_title = (doc.get("title") or "document").replace(" ", "_")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{safe_title}.docx"'},
    )


# ── Delete ────────────────────────────────────────────────────────────────────

@router.delete("/{document_id}", status_code=status.HTTP_200_OK)
async def remove_document(
    document_id: str,
    request: Request,
    db: AsyncSession = Depends(get_db),
) -> None:
    """
    Deletes document from Postgres.
    Knowledge units are CASCADE deleted.
    Emits dataplane.documents.deleted — Embedding Worker cleans up Qdrant vectors.
    """
    auth = _get_auth(request)
    access = _get_org_access(request)

    if not access.has_permission("resources:delete"):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=f"Permission 'resources:delete' required. Your role '{access.role}' "
                   "does not have this. Only admins and owners can delete documents.",
        )

    deleted = await delete_document(db, document_id=document_id, org_id=auth.org_id)
    if not deleted:
        raise HTTPException(status_code=404, detail="Document not found")
    await publish_document_deleted(document_id=document_id, org_id=auth.org_id)


# ── Internal status update (called by workers) ────────────────────────────────

@router.patch("/{document_id}/status", status_code=status.HTTP_200_OK)
async def patch_status(
    document_id: str,
    new_status: str = Query(..., alias="status"),
    error: Optional[str] = Query(None),
    db: AsyncSession = Depends(get_db),
    _: bool = Depends(_verify_internal_key),
) -> None:
    await update_document_status(
        db, document_id=document_id, status=new_status, error_message=error
    )
