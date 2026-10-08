import base64
import io
import zipfile
from pathlib import Path
import pytest
from docx import Document
from fastapi import HTTPException
from app.chat_extract import ChatExtractRequest, extract_chat_document


def request(name, content):
    return ChatExtractRequest(filename=name, content_base64=base64.b64encode(content).decode())


# Hand-built minimal OOXML zips: no openpyxl/python-pptx dependency exists in
# this service (nor does app.parsers use one), so the fixtures below assemble
# just the zip parts app.parsers._parse_xlsx / _parse_pptx actually read.

class _Cell:
    """A worksheet cell written the way a real producer stores it: an explicit
    type (`t`), a style index (`s`), an optional `<f>` and an optional `<v>`.

    `v=None` writes no `<v>` element at all; `v=''` writes an empty `<v></v>`
    (what openpyxl writes for a formula it did not calculate)."""

    def __init__(self, v=None, *, t=None, s=None, f=None, f_attrs=''):
        self.v, self.t, self.s, self.f, self.f_attrs = v, t, s, f, f_attrs

    def xml(self, ref: str) -> str:
        attrs = f' r="{ref}"'
        if self.t:
            attrs += f' t="{self.t}"'
        if self.s is not None:
            attrs += f' s="{self.s}"'
        formula = ''
        if self.f is not None or self.f_attrs:
            formula = f'<f{self.f_attrs}>{self.f or ""}</f>' if self.f else f'<f{self.f_attrs}/>'
        value = '' if self.v is None else f'<v>{self.v}</v>'
        return f'<c{attrs}>{formula}{value}</c>'


_STYLES_NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'


def _styles_xml(xfs: list[int], custom_formats: dict[int, str] | None = None) -> str:
    """A minimal xl/styles.xml whose cellXfs index `i` uses number format `xfs[i]`."""
    num_fmts = ''.join(
        f'<numFmt numFmtId="{i}" formatCode="{code}"/>' for i, code in (custom_formats or {}).items()
    )
    return (
        '<?xml version="1.0"?>'
        f'<styleSheet xmlns="{_STYLES_NS}">'
        + (f'<numFmts count="{len(custom_formats)}">{num_fmts}</numFmts>' if custom_formats else '')
        + f'<cellXfs count="{len(xfs)}">'
        + ''.join(f'<xf numFmtId="{n}" applyNumberFormat="1"/>' for n in xfs)
        + '</cellXfs></styleSheet>'
    )


def _build_xlsx(
    sheets: list[tuple[str, list[list]]],
    *,
    styles_xml: str | None = None,
    date1904: bool = False,
    states: dict[str, str] | None = None,
) -> bytes:
    shared: list[str] = []

    def shared_index(text: str) -> int:
        if text not in shared:
            shared.append(text)
        return shared.index(text)

    sheet_entries = []
    worksheet_files = {}
    for sheet_index, (name, rows) in enumerate(sheets, start=1):
        row_xml = []
        for row_index, row in enumerate(rows, start=1):
            if row is None:  # a row number the producer never wrote
                continue
            cells_xml = []
            for col_index, value in enumerate(row):
                col_letter = chr(ord('A') + col_index)
                ref = f'{col_letter}{row_index}'
                if isinstance(value, _Cell):
                    cells_xml.append(value.xml(ref))
                elif value is None:
                    continue
                else:
                    cells_xml.append(f'<c r="{ref}" t="s"><v>{shared_index(value)}</v></c>')
            row_xml.append(f'<row r="{row_index}">{"".join(cells_xml)}</row>')
        worksheet_files[f'xl/worksheets/sheet{sheet_index}.xml'] = (
            '<?xml version="1.0"?>'
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            f'<sheetData>{"".join(row_xml)}</sheetData></worksheet>'
        )
        state = f' state="{states[name]}"' if states and name in states else ''
        sheet_entries.append(
            f'<sheet name="{name}" sheetId="{sheet_index}"{state} r:id="rId{sheet_index}"/>'
        )

    workbook_xml = (
        '<?xml version="1.0"?>'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        + ('<workbookPr date1904="1"/>' if date1904 else '<workbookPr/>')
        + f'<sheets>{"".join(sheet_entries)}</sheets></workbook>'
    )
    rels_xml = (
        '<?xml version="1.0"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + ''.join(
            f'<Relationship Id="rId{i}" Target="worksheets/sheet{i}.xml"/>'
            for i in range(1, len(sheets) + 1)
        )
        + '</Relationships>'
    )
    shared_strings_xml = (
        '<?xml version="1.0"?>'
        f'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="{len(shared)}" uniqueCount="{len(shared)}">'
        + ''.join(f'<si><t>{text}</t></si>' for text in shared)
        + '</sst>'
    )

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w') as zf:
        zf.writestr('xl/workbook.xml', workbook_xml)
        zf.writestr('xl/_rels/workbook.xml.rels', rels_xml)
        zf.writestr('xl/sharedStrings.xml', shared_strings_xml)
        if styles_xml is not None:
            zf.writestr('xl/styles.xml', styles_xml)
        for path, xml in worksheet_files.items():
            zf.writestr(path, xml)
    return buffer.getvalue()


def _build_pptx(slides: list[list[str]]) -> bytes:
    sld_id_list = []
    rels = []
    slide_files = {}
    for index, paragraphs in enumerate(slides, start=1):
        rid = f'rId{index + 1}'
        sld_id_list.append(f'<p:sldId id="{255 + index}" r:id="{rid}"/>')
        rels.append(f'<Relationship Id="{rid}" Target="slides/slide{index}.xml"/>')
        paragraph_xml = ''.join(
            f'<a:p><a:r><a:t>{text}</a:t></a:r></a:p>' for text in paragraphs
        )
        slide_files[f'ppt/slides/slide{index}.xml'] = (
            '<?xml version="1.0"?>'
            '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            f'<p:cSld><p:spTree><p:sp><p:txBody>{paragraph_xml}</p:txBody></p:sp></p:spTree></p:cSld></p:sld>'
        )

    presentation_xml = (
        '<?xml version="1.0"?>'
        '<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'<p:sldIdLst>{"".join(sld_id_list)}</p:sldIdLst></p:presentation>'
    )
    rels_xml = (
        '<?xml version="1.0"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + ''.join(rels)
        + '</Relationships>'
    )

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w') as zf:
        zf.writestr('ppt/presentation.xml', presentation_xml)
        zf.writestr('ppt/_rels/presentation.xml.rels', rels_xml)
        for path, xml in slide_files.items():
            zf.writestr(path, xml)
    return buffer.getvalue()


_P_NS = 'http://schemas.openxmlformats.org/presentationml/2006/main'
_A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'
_C_NS = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
_R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
_REL_TYPE = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/'


def _para(text: str) -> str:
    return f'<a:p><a:r><a:t>{text}</a:t></a:r></a:p>'


def _shape(name: str, paragraphs: list[str], ph: str = '') -> str:
    placeholder = f'<p:ph {ph}/>' if ph else ''
    return (
        f'<p:sp><p:nvSpPr><p:cNvPr id="2" name="{name}"/><p:cNvSpPr/><p:nvPr>{placeholder}</p:nvPr></p:nvSpPr>'
        f'<p:spPr/><p:txBody><a:bodyPr/>{"".join(_para(t) for t in paragraphs)}</p:txBody></p:sp>'
    )


def _title_shape(text: str) -> str:
    return _shape('Title 1', [text], 'type="title"')


def _picture(descr: str) -> str:
    return (
        f'<p:pic><p:nvPicPr><p:cNvPr id="5" name="Picture 4" descr="{descr}"/><p:cNvPicPr/><p:nvPr/></p:nvPicPr>'
        '<p:blipFill><a:blip r:embed="rId9"/></p:blipFill><p:spPr/></p:pic>'
    )


def _graphic_frame(graphic_data: str) -> str:
    return (
        '<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="4" name="Frame 3"/><p:cNvGraphicFramePr/><p:nvPr/>'
        '</p:nvGraphicFramePr><p:xfrm><a:off x="0" y="0"/><a:ext cx="1" cy="1"/></p:xfrm>'
        f'<a:graphic>{graphic_data}</a:graphic></p:graphicFrame>'
    )


def _table_frame(rows: list[list[str]]) -> str:
    trs = ''.join(
        '<a:tr h="370840">'
        + ''.join(f'<a:tc><a:txBody><a:bodyPr/>{_para(cell)}</a:txBody></a:tc>' for cell in row)
        + '</a:tr>'
        for row in rows
    )
    return _graphic_frame(
        '<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/table">'
        f'<a:tbl><a:tblGrid/>{trs}</a:tbl></a:graphicData>'
    )


def _chart_frame(rid: str) -> str:
    return _graphic_frame(
        '<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
        f'<c:chart xmlns:c="{_C_NS}" r:id="{rid}"/></a:graphicData>'
    )


def _chart_xml(title: str | None, series: list[tuple[str, list[str]]], categories: list[str]) -> str:
    cat_pts = ''.join(f'<c:pt idx="{i}"><c:v>{c}</c:v></c:pt>' for i, c in enumerate(categories))
    ser_xml = ''
    for i, (name, values) in enumerate(series):
        val_pts = ''.join(f'<c:pt idx="{j}"><c:v>{v}</c:v></c:pt>' for j, v in enumerate(values))
        ser_xml += (
            f'<c:ser><c:idx val="{i}"/><c:order val="{i}"/>'
            '<c:tx><c:strRef><c:f>Sheet1!$B$1</c:f><c:strCache><c:ptCount val="1"/>'
            f'<c:pt idx="0"><c:v>{name}</c:v></c:pt></c:strCache></c:strRef></c:tx>'
            '<c:cat><c:strRef><c:f>Sheet1!$A$2:$A$4</c:f>'
            f'<c:strCache><c:ptCount val="{len(categories)}"/>{cat_pts}</c:strCache></c:strRef></c:cat>'
            '<c:val><c:numRef><c:f>Sheet1!$B$2:$B$4</c:f><c:numCache><c:formatCode>General</c:formatCode>'
            f'<c:ptCount val="{len(values)}"/>{val_pts}</c:numCache></c:numRef></c:val></c:ser>'
        )
    title_xml = (
        '<c:title><c:tx><c:rich><a:bodyPr/>'
        f'<a:p><a:r><a:t>{title}</a:t></a:r></a:p></c:rich></c:tx><c:overlay val="0"/></c:title>'
        if title else ''
    )
    return (
        '<?xml version="1.0"?>'
        f'<c:chartSpace xmlns:c="{_C_NS}" xmlns:a="{_A_NS}"><c:date1904 val="0"/>'
        f'<c:chart>{title_xml}<c:plotArea><c:barChart><c:barDir val="col"/>{ser_xml}</c:barChart></c:plotArea>'
        '</c:chart></c:chartSpace>'
    )


def _notes_xml(text: str) -> str:
    return (
        '<?xml version="1.0"?>'
        f'<p:notes xmlns:p="{_P_NS}" xmlns:a="{_A_NS}"><p:cSld><p:spTree>'
        '<p:sp><p:nvSpPr><p:cNvPr id="2" name="Slide Image Placeholder 1"/><p:cNvSpPr/>'
        '<p:nvPr><p:ph type="sldImg"/></p:nvPr></p:nvSpPr><p:spPr/></p:sp>'
        + _shape('Notes Placeholder 2', [text], 'type="body" idx="1"')
        + '<p:sp><p:nvSpPr><p:cNvPr id="4" name="Slide Number Placeholder 3"/><p:cNvSpPr/>'
        '<p:nvPr><p:ph type="sldNum" sz="quarter" idx="5"/></p:nvPr></p:nvSpPr><p:spPr/>'
        '<p:txBody><a:bodyPr/><a:p><a:fld id="{00000000-0000-0000-0000-000000000000}" type="slidenum">'
        '<a:t>777</a:t></a:fld></a:p></p:txBody></p:sp>'
        '</p:spTree></p:cSld></p:notes>'
    )


def _build_pptx_rich(slides: list[dict]) -> bytes:
    """Slides shaped like PowerPoint's own output. Each slide is a dict with
    `shapes` (spTree children as XML), `show` (False = hidden), `notes` (text
    for a notesSlide) and `charts` ({relationship id: chart part XML})."""
    sld_ids, presentation_rels, parts = [], [], {}
    for index, slide in enumerate(slides, start=1):
        rid = f'rId{index + 1}'
        sld_ids.append(f'<p:sldId id="{255 + index}" r:id="{rid}"/>')
        presentation_rels.append(
            f'<Relationship Id="{rid}" Type="{_REL_TYPE}slide" Target="slides/slide{index}.xml"/>'
        )
        show = '' if slide.get('show', True) else ' show="0"'
        parts[f'ppt/slides/slide{index}.xml'] = (
            '<?xml version="1.0"?>'
            f'<p:sld xmlns:p="{_P_NS}" xmlns:a="{_A_NS}" xmlns:r="{_R_NS}"{show}>'
            f'<p:cSld><p:spTree>{slide["shapes"]}</p:spTree></p:cSld></p:sld>'
        )
        rels = [f'<Relationship Id="rId1" Type="{_REL_TYPE}slideLayout" Target="../slideLayouts/slideLayout1.xml"/>']
        if slide.get('notes') is not None:
            rels.append(
                f'<Relationship Id="rId2" Type="{_REL_TYPE}notesSlide" Target="../notesSlides/notesSlide{index}.xml"/>'
            )
            parts[f'ppt/notesSlides/notesSlide{index}.xml'] = _notes_xml(slide['notes'])
        for chart_rid, chart in slide.get('charts', {}).items():
            rels.append(
                f'<Relationship Id="{chart_rid}" Type="{_REL_TYPE}chart" Target="../charts/chart{index}.xml"/>'
            )
            parts[f'ppt/charts/chart{index}.xml'] = chart
        parts[f'ppt/slides/_rels/slide{index}.xml.rels'] = (
            '<?xml version="1.0"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            + ''.join(rels) + '</Relationships>'
        )
    parts['ppt/presentation.xml'] = (
        '<?xml version="1.0"?>'
        f'<p:presentation xmlns:p="{_P_NS}" xmlns:r="{_R_NS}">'
        f'<p:sldIdLst>{"".join(sld_ids)}</p:sldIdLst></p:presentation>'
    )
    parts['ppt/_rels/presentation.xml.rels'] = (
        '<?xml version="1.0"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + ''.join(presentation_rels) + '</Relationships>'
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w') as zf:
        for path, xml in parts.items():
            zf.writestr(path, xml)
    return buffer.getvalue()


@pytest.mark.parametrize('name', ['source.md', 'sales.csv'])
def test_text_extraction_is_complete_and_ephemeral(name):
    result = extract_chat_document(request(name, b'first\nlast'))
    assert result['data']['content'] == 'first\nlast'
    assert result['data']['persisted'] is False


def test_docx_extraction():
    document = Document()
    document.add_paragraph('Order FF-1042: 9 packed, 3 pending.')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Current policy'
    table.cell(0, 1).text = 'K3'
    document.add_paragraph('Last paragraph')
    buffer = io.BytesIO()
    document.save(buffer)
    text = extract_chat_document(request('order.docx', buffer.getvalue()))['data']['content']
    assert '9 packed, 3 pending' in text
    assert text.index('9 packed') < text.index('Current policy\tK3') < text.index('Last paragraph')


def test_xlsx_extraction_keeps_sheet_order_and_cell_values():
    content = _build_xlsx([
        ('Orders', [['Order', 'Qty'], ['FF-1042', '9']]),
        ('Notes', [['Last row']]),
    ])
    text = extract_chat_document(request('report.xlsx', content))['data']['content']
    assert text.index('[Orders]') < text.index('FF-1042\t9') < text.index('[Notes]') < text.index('Last row')


def test_pptx_extraction_keeps_slide_order_and_text():
    content = _build_pptx([
        ['Title One', 'Bullet A'],
        ['Title Two'],
    ])
    text = extract_chat_document(request('deck.pptx', content))['data']['content']
    assert text.index('[Slide 1]') < text.index('Title One') < text.index('Bullet A') < text.index('[Slide 2]') < text.index('Title Two')


@pytest.mark.parametrize('name,content,status', [
    ('broken.pdf', b'not a PDF', 422), ('empty.md', b' ', 422),
    ('large.md', b'x' * 60_001, 413), ('unknown.exe', b'binary', 415),
    ('broken.xlsx', b'not a zip', 422), ('broken.pptx', b'not a zip', 422),
], ids=['broken-pdf', 'empty-md', 'large-md', 'unknown-exe', 'broken-xlsx', 'broken-pptx'])
def test_unreadable_or_oversized_input_is_rejected(name, content, status):
    with pytest.raises(HTTPException) as exc:
        extract_chat_document(request(name, content))
    assert exc.value.status_code == status


# OPC relationship targets are normally relative to the referencing part's own
# folder ("worksheets/sheet1.xml", "slides/slide1.xml"), but the spec also
# allows a package-root-absolute form ("/xl/worksheets/sheet1.xml"). Real
# Excel/PowerPoint output never uses the absolute form, but nothing rules it
# out, and _xlsx_sheet_parts/_pptx_slide_parts must not silently read the
# sheet/slide as empty when a producer does.

def test_xlsx_with_absolute_style_relationship_target_still_reads_sheet():
    workbook_xml = (
        '<?xml version="1.0"?>'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<sheets><sheet name="Orders" sheetId="1" r:id="rId1"/></sheets></workbook>'
    )
    rels_xml = (
        '<?xml version="1.0"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Target="/xl/worksheets/sheet1.xml"/></Relationships>'
    )
    sheet_xml = (
        '<?xml version="1.0"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<sheetData><row r="1"><c r="A1"><v>9</v></c></row></sheetData></worksheet>'
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w') as zf:
        zf.writestr('xl/workbook.xml', workbook_xml)
        zf.writestr('xl/_rels/workbook.xml.rels', rels_xml)
        zf.writestr('xl/worksheets/sheet1.xml', sheet_xml)
    text = extract_chat_document(request('report.xlsx', buffer.getvalue()))['data']['content']
    assert text == '[Orders]\n9'


def test_pptx_with_absolute_style_relationship_target_still_reads_slide():
    presentation_xml = (
        '<?xml version="1.0"?>'
        '<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<p:sldIdLst><p:sldId id="256" r:id="rId1"/></p:sldIdLst></p:presentation>'
    )
    rels_xml = (
        '<?xml version="1.0"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Target="/ppt/slides/slide1.xml"/></Relationships>'
    )
    slide_xml = (
        '<?xml version="1.0"?>'
        '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:t>Title</a:t></a:r></a:p>'
        '</p:txBody></p:sp></p:spTree></p:cSld></p:sld>'
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w') as zf:
        zf.writestr('ppt/presentation.xml', presentation_xml)
        zf.writestr('ppt/_rels/presentation.xml.rels', rels_xml)
        zf.writestr('ppt/slides/slide1.xml', slide_xml)
    text = extract_chat_document(request('deck.pptx', buffer.getvalue()))['data']['content']
    assert text == '[Slide 1]\nTitle'


# ---------------------------------------------------------------------------
# H1: read what Excel and PowerPoint actually store.

def read(name: str, content: bytes) -> str:
    return extract_chat_document(request(name, content))['data']['content']


def one_sheet(rows, **kwargs) -> str:
    return read('book.xlsx', _build_xlsx([('Ark', rows)], **kwargs))


def test_xlsx_formula_without_a_cached_value_prints_the_formula():
    text = one_sheet([['Post', 'Beløp'], ['a', _Cell('4')], ['Sum', _Cell(None, f='SUM(B2:B3)')]])
    assert 'Sum\t=SUM(B2:B3)' in text
    assert 'Ark!B3: =SUM(B2:B3)' in text


def test_xlsx_formula_with_an_empty_value_element_prints_the_formula():
    # openpyxl writes <f>..</f><v></v> for a formula it never calculated.
    text = one_sheet([['Sum', _Cell('', f='SUM(B2:B9)')]])
    assert text.startswith('[Ark]\nSum\t=SUM(B2:B9)')


def test_xlsx_formula_with_a_cached_value_prints_the_value_and_lists_the_formula():
    text = one_sheet([['Sum', _Cell('10', f='SUM(B2:B3)')]])
    grid, _, formulas = text.partition('\n\n[Formler]\n')
    assert grid == '[Ark]\nSum\t10'
    assert formulas == 'Ark!B1: =SUM(B2:B3)'


def test_xlsx_string_formula_with_an_empty_result_stays_empty():
    # Excel writes <c t="str"><f>..</f><v></v></c> for a formula that returns "".
    text = one_sheet([['x', _Cell('', t='str', f='IF(A1="","","y")')]])
    assert text.startswith('[Ark]\nx\n')


def test_xlsx_shared_formula_child_prints_a_placeholder_and_the_master_is_listed_once():
    text = one_sheet([
        [_Cell('2', f='A2*2', f_attrs=' t="shared" ref="A1:A2" si="0"')],
        [_Cell(None, f_attrs=' t="shared" si="0"')],
    ])
    assert '=(delt formel)' in text
    formulas = text.partition('[Formler]\n')[2]
    assert formulas.count('A2*2') == 1
    assert 'A1:A2' in formulas


def test_xlsx_formula_list_quotes_sheet_names_with_spaces():
    text = read('book.xlsx', _build_xlsx([('Mitt ark', [[_Cell('2', f='1+1')]])]))
    assert "'Mitt ark'!A1: =1+1" in text


def test_xlsx_formula_list_is_capped_at_100_lines():
    rows = [[_Cell(str(i), f=f'{i}+0')] for i in range(150)]
    text = one_sheet(rows)
    formulas = text.partition('[Formler]\n')[2].split('\n')
    assert sum(1 for line in formulas if line.startswith('Ark!A')) == 100
    assert formulas[-1] == '… 50 flere formler utelatt'


def test_xlsx_serial_in_a_date_style_prints_an_iso_date():
    styles = _styles_xml([0, 14, 164, 165, 20], {164: r'd\.m\.yyyy', 165: r'dd\/mm\/yyyy\ hh:mm'})
    text = one_sheet([[
        _Cell('45292', s=1), _Cell('45366', s=2), _Cell('45292.5', s=3), _Cell('0.75', s=4), _Cell('45292', s=0),
    ]], styles_xml=styles)
    assert text == '[Ark]\n2024-01-01\t2024-03-15\t2024-01-01 12:00\t18:00\t45292'


def test_xlsx_date1904_workbook_uses_the_1904_epoch():
    styles = _styles_xml([0, 14])
    text = one_sheet([[_Cell('43830', s=1)]], styles_xml=styles, date1904=True)
    assert text == '[Ark]\n2024-01-01'
    # The same serial in a 1900 workbook is a different day (1904 offset is 1462 days).
    assert one_sheet([[_Cell('43830', s=1)]], styles_xml=styles) == '[Ark]\n2019-12-31'


def test_excel_serial_conversion_keeps_the_1900_leap_day_quirk():
    from app.parsers import _excel_serial_to_text
    assert _excel_serial_to_text(1, 'date') == '1900-01-01'
    assert _excel_serial_to_text(59, 'date') == '1900-02-28'
    assert _excel_serial_to_text(60, 'date') == '1900-02-29'  # the day Excel invented
    assert _excel_serial_to_text(61, 'date') == '1900-03-01'
    assert _excel_serial_to_text(45292, 'date') == '2024-01-01'
    assert _excel_serial_to_text(45292.999999, 'datetime') == '2024-01-02 00:00'  # rounds up, no 23:59:60
    assert _excel_serial_to_text(-5, 'date') is None
    assert _excel_serial_to_text(9_999_999, 'date') is None


def test_xlsx_custom_formats_are_classified_by_their_code():
    from app.parsers import _classify_number_format
    assert _classify_number_format(r'd\.m\.yyyy').kind == 'date'
    assert _classify_number_format('[$-409]mmmm d, yyyy;@').kind == 'date'
    assert _classify_number_format('yyyy-mm-dd hh:mm').kind == 'datetime'
    assert _classify_number_format('h:mm:ss').kind == 'time'
    assert _classify_number_format('[h]:mm').kind == 'time'
    assert _classify_number_format('mm:ss').kind == 'time'   # minutes, not months
    assert _classify_number_format('0.0%').kind == 'percent'
    assert _classify_number_format('0.00%').decimals == 2
    # Literals and colours must not turn a number format into a date format.
    assert _classify_number_format('0.00 "dager"').kind == 'general'
    assert _classify_number_format('General" kr"').kind == 'general'
    assert _classify_number_format('[Red]#,##0;[Blue]-#,##0').kind == 'general'
    assert _classify_number_format('0 "%"').kind == 'general'


def test_xlsx_percent_styles_print_percentages_without_float_noise():
    styles = _styles_xml([0, 9, 10, 166], {166: '0.0%'})
    text = one_sheet([[
        _Cell('0.255', s=3), _Cell('0.3', s=2), _Cell('0.30000000000000004', s=1), _Cell('0.5', s=0),
    ]], styles_xml=styles)
    assert text == '[Ark]\n25.5%\t30.00%\t30%\t0.5'


def test_xlsx_booleans_and_errors():
    text = one_sheet([[_Cell('1', t='b'), _Cell('0', t='b'), _Cell('#DIV/0!', t='e'), _Cell('#N/A', t='e', f='NA()')]])
    assert text.startswith('[Ark]\nTRUE\tFALSE\t#DIV/0!\t#N/A')


def test_xlsx_numbers_lose_their_float_noise_but_keep_integers():
    text = one_sheet([[
        _Cell('0.30000000000000004'), _Cell('12.5'), _Cell('100'), _Cell('1.0000000000000002'),
        _Cell('12345678901234567'), _Cell('1E-07'), _Cell('abc', t='str'),
    ]])
    assert text == '[Ark]\n0.3\t12.5\t100\t1\t12345678901234567\t0.0000001\tabc'


def test_xlsx_hidden_sheets_are_labelled():
    content = _build_xlsx(
        [('Synlig', [['a']]), ('Skjult', [['b']]), ('Helt skjult', [['c']])],
        states={'Skjult': 'hidden', 'Helt skjult': 'veryHidden'},
    )
    text = read('book.xlsx', content)
    assert '[Synlig]\na' in text
    assert '[Skjult] (skjult)\nb' in text
    assert '[Helt skjult] (skjult)\nc' in text


def test_xlsx_rows_keep_their_position_across_gaps_and_empty_rows():
    # Row n is line n of the sheet, so "what is in A3" has one answer.
    text = one_sheet([['Head'], None, None, ['After gap'], [], ['Last']])
    assert text == '[Ark]\nHead\n\n\nAfter gap\n\nLast'
    assert one_sheet([None, None, ['Third row']]) == '[Ark]\n\n\nThird row'


def test_xlsx_a_long_run_of_missing_rows_prints_a_jump_marker():
    rows = [['top']] + [None] * 200 + [['bottom']]
    assert one_sheet(rows) == '[Ark]\ntop\n[hopper til rad 202]\nbottom'


def test_xlsx_trailing_formatted_empty_cells_leave_no_tabs():
    text = one_sheet([['x', _Cell(None, s=1), _Cell(None, s=1)], ['y']], styles_xml=_styles_xml([0, 14]))
    assert text == '[Ark]\nx\ny'


def _long_sheet(name: str, rows: int) -> tuple[str, list[list[str]]]:
    return name, [[f'Rad {i:05d}', '12345'] for i in range(rows)]


def test_xlsx_over_the_cap_is_cut_on_a_row_boundary_instead_of_refused():
    import re
    result = extract_chat_document(request('big.xlsx', _build_xlsx([_long_sheet('Data', 6000)])))['data']
    assert result['truncated'] is True
    lines = result['content'].split('\n')
    assert lines[0] == '[Data]'
    marker = re.fullmatch(r'\[… avkortet: (\d+) rader utelatt i «Data»\]', lines[-1])
    assert marker, lines[-1]
    kept = lines[1:-1]
    assert all(re.fullmatch(r'Rad \d{5}\t12345', line) for line in kept)  # no half-written row
    assert len(kept) + int(marker.group(1)) == 6000
    assert 50_000 < len(result['content']) <= 60_000


def test_xlsx_cut_marks_the_sheets_that_were_left_out():
    content = _build_xlsx([_long_sheet('Data', 6000), ('Notater', [['liten']])])
    text = read('big.xlsx', content)
    assert '[… avkortet: 1 rader utelatt i «Notater»]' in text
    assert '[Notater]' not in text


def test_xlsx_under_the_cap_is_not_marked_truncated():
    result = extract_chat_document(request('small.xlsx', _build_xlsx([('Ark', [['a', 'b']])])))['data']
    assert 'truncated' not in result


def test_xlsx_formula_list_after_a_cut_still_fits_the_turn_limit():
    rows = [['x' * 100, _Cell(str(i), f=f'SUM(A1:A{i + 1})+' + 'B1+' * 40 + '0')] for i in range(700)]
    result = extract_chat_document(request('big.xlsx', _build_xlsx([('Data', rows)])))['data']
    assert result['truncated'] is True
    assert len(result['content']) <= 60_000
    assert '[Formler]' in result['content']


# --- PowerPoint -----------------------------------------------------------

def test_pptx_speaker_notes_come_after_the_slide_text_and_only_the_body_is_read():
    content = _build_pptx_rich([
        {'shapes': _title_shape('Forside')},
        {'shapes': _title_shape('Resultater'), 'notes': 'Husk å nevne den nye lederen.'},
    ])
    text = read('deck.pptx', content)
    assert text == '[Slide 1]\nForside\n\n[Slide 2]\nResultater\n[Notater]\nHusk å nevne den nye lederen.'
    assert '777' not in text  # the notes page's slide-number field


def test_pptx_empty_speaker_notes_add_no_block():
    text = read('deck.pptx', _build_pptx_rich([{'shapes': _title_shape('A'), 'notes': '  '}]))
    assert text == '[Slide 1]\nA'


def test_pptx_tables_become_tab_separated_rows_in_slide_order():
    content = _build_pptx_rich([{'shapes': (
        _title_shape('Tall')
        + _shape('Body', ['Før tabellen'])
        + _table_frame([['Region', 'Omsetning'], ['Nord', '120'], ['Sør', '95']])
        + _shape('Footer', ['Etter tabellen'])
    )}])
    text = read('deck.pptx', content)
    assert text == '[Slide 1]\nTall\nFør tabellen\nRegion\tOmsetning\nNord\t120\nSør\t95\nEtter tabellen'


def test_pptx_the_title_comes_first_wherever_it_sits_in_the_shape_tree():
    content = _build_pptx_rich([{'shapes': _shape('Body', ['Brødtekst']) + _title_shape('Tittelen')}])
    assert read('deck.pptx', content) == '[Slide 1]\nTittelen\nBrødtekst'


def test_pptx_chart_values_are_read_from_the_cached_points():
    chart = _chart_xml('Omsetning', [('Salg', ['10', '20', '30'])], ['Q1', 'Q2', 'Q3'])
    content = _build_pptx_rich([{
        'shapes': _title_shape('Utvikling') + _chart_frame('rId5'),
        'charts': {'rId5': chart},
    }])
    assert read('deck.pptx', content) == (
        '[Slide 1]\nUtvikling\n[Diagram: Omsetning] Salg: 10, 20, 30; kategorier: Q1, Q2, Q3'
    )


def test_pptx_chart_without_a_title_and_with_two_series():
    chart = _chart_xml(None, [('A', ['1', '2']), ('B', ['0.30000000000000004', '4'])], ['x', 'y'])
    content = _build_pptx_rich([{'shapes': _chart_frame('rId5'), 'charts': {'rId5': chart}}])
    assert read('deck.pptx', content) == '[Slide 1]\n[Diagram] A: 1, 2; B: 0.3, 4; kategorier: x, y'


def _num_cache(values: list, format_code: str = 'General', count: int | None = None) -> str:
    pts = ''.join(f'<c:pt idx="{i}"><c:v>{v}</c:v></c:pt>' for i, v in values if v is not None)
    return (f'<c:numCache><c:formatCode>{format_code}</c:formatCode>'
            f'<c:ptCount val="{count if count is not None else len(values)}"/>{pts}</c:numCache>')


def _str_cache(values: list[str]) -> str:
    pts = ''.join(f'<c:pt idx="{i}"><c:v>{v}</c:v></c:pt>' for i, v in enumerate(values))
    return f'<c:strCache><c:ptCount val="{len(values)}"/>{pts}</c:strCache>'


def _chart_deck_text(plot_area: str) -> str:
    chart = (
        '<?xml version="1.0"?>'
        f'<c:chartSpace xmlns:c="{_C_NS}" xmlns:a="{_A_NS}"><c:date1904 val="0"/>'
        f'<c:chart><c:plotArea>{plot_area}</c:plotArea></c:chart></c:chartSpace>'
    )
    content = _build_pptx_rich([{'shapes': _chart_frame('rId5'), 'charts': {'rId5': chart}}])
    return read('deck.pptx', content)


def _series(name: str, *, cat: str = '', val: str = '', x: str = '', y: str = '') -> str:
    return (
        f'<c:ser><c:idx val="0"/><c:order val="0"/><c:tx><c:strRef><c:f>S!$B$1</c:f>{_str_cache([name])}</c:strRef></c:tx>'
        + (f'<c:cat>{cat}</c:cat>' if cat else '') + (f'<c:val>{val}</c:val>' if val else '')
        + (f'<c:xVal>{x}</c:xVal>' if x else '') + (f'<c:yVal>{y}</c:yVal>' if y else '') + '</c:ser>'
    )


def test_pptx_scatter_chart_reads_x_and_y_points():
    plot = '<c:scatterChart>' + _series(
        'Y', x=f'<c:numRef><c:f>S!$A$2:$A$3</c:f>{_num_cache([(0, 1), (1, 2)])}</c:numRef>',
        y=f'<c:numRef><c:f>S!$B$2:$B$3</c:f>{_num_cache([(0, 10), (1, 20)])}</c:numRef>',
    ) + '</c:scatterChart>'
    assert _chart_deck_text(plot) == '[Slide 1]\n[Diagram] Y: 10, 20; kategorier: 1, 2'


def test_pptx_chart_dates_and_percentages_follow_the_cached_format_code():
    plot = '<c:lineChart>' + _series(
        'Andel',
        cat=f'<c:numRef><c:f>S!$A$2:$A$3</c:f>{_num_cache([(0, 45292), (1, 45293)], r"d\.m\.yyyy")}</c:numRef>',
        val=f'<c:numRef><c:f>S!$B$2:$B$3</c:f>{_num_cache([(0, 0.25), (1, 0.5)], "0%")}</c:numRef>',
    ) + '</c:lineChart>'
    assert _chart_deck_text(plot) == '[Slide 1]\n[Diagram] Andel: 25%, 50%; kategorier: 2024-01-01, 2024-01-02'


def test_pptx_chart_with_multi_level_categories_and_a_gap_in_the_points():
    levels = (
        '<c:multiLvlStrRef><c:f>S!$A$2:$B$3</c:f><c:multiLvlStrCache><c:ptCount val="2"/>'
        '<c:lvl><c:pt idx="0"><c:v>Jan</c:v></c:pt><c:pt idx="1"><c:v>Feb</c:v></c:pt></c:lvl>'
        '<c:lvl><c:pt idx="0"><c:v>2024</c:v></c:pt></c:lvl></c:multiLvlStrCache></c:multiLvlStrRef>'
    )
    gap = f'<c:numRef><c:f>S!$C$2:$C$4</c:f>{_num_cache([(0, 1), (2, 3)], count=3)}</c:numRef>'
    plot = '<c:barChart>' + _series('Verdi', cat=levels, val=gap) + '</c:barChart>'
    assert _chart_deck_text(plot) == '[Slide 1]\n[Diagram] Verdi: 1, , 3; kategorier: Jan, Feb'


def test_pptx_chart_points_are_capped_with_a_marker():
    values = [(i, i) for i in range(150)]
    plot = '<c:barChart>' + _series('Mange', val=f'<c:numRef><c:f>S!$B$2:$B$151</c:f>{_num_cache(values)}</c:numRef>') + '</c:barChart>'
    text = _chart_deck_text(plot)
    assert text.startswith('[Slide 1]\n[Diagram] Mange: 0, 1, 2,')
    assert text.endswith('99, … (+50)')


def test_pptx_hidden_slides_are_labelled():
    content = _build_pptx_rich([
        {'shapes': _title_shape('Synlig')},
        {'shapes': _title_shape('Reserve'), 'show': False},
    ])
    assert read('deck.pptx', content) == '[Slide 1]\nSynlig\n\n[Slide 2] (skjult)\nReserve'


def test_pptx_picture_alt_text_is_kept_and_groups_are_walked_in_order():
    group = f'<p:grpSp><p:nvGrpSpPr><p:cNvPr id="7" name="Group 6"/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/>{_shape("In group", ["I gruppen"])}</p:grpSp>'
    content = _build_pptx_rich([{'shapes': _shape('A', ['Først']) + _picture('Firmalogo') + group + _shape('B', ['Sist'])}])
    assert read('deck.pptx', content) == '[Slide 1]\nFørst\n[Bilde: Firmalogo]\nI gruppen\nSist'


def test_pptx_line_breaks_inside_a_paragraph_are_kept():
    shape = (
        '<p:sp><p:nvSpPr><p:cNvPr id="2" name="B"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/>'
        '<a:p><a:r><a:t>Linje 1</a:t></a:r><a:br/><a:r><a:t>Linje 2</a:t></a:r></a:p></p:txBody></p:sp>'
    )
    assert read('deck.pptx', _build_pptx_rich([{'shapes': shape}])) == '[Slide 1]\nLinje 1\nLinje 2'


# --- Zip budget -----------------------------------------------------------

def _zeros_zip(sizes: dict[str, int], files: dict[str, str] | None = None) -> bytes:
    """A real zip whose members expand to `sizes` bytes of zeros: a few hundred
    KB on the wire, hundreds of MB if someone reads it all into memory."""
    buffer = io.BytesIO()
    chunk = b'\0' * (1 << 20)
    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED, compresslevel=9) as zf:
        for path, xml in (files or {}).items():
            zf.writestr(path, xml)
        for path, size in sizes.items():
            with zf.open(path, 'w', force_zip64=True) as part:
                for _ in range(size // len(chunk)):
                    part.write(chunk)
    return buffer.getvalue()


_MINI_WORKBOOK = {
    'xl/workbook.xml': (
        '<?xml version="1.0"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<sheets><sheet name="Ark" sheetId="1" r:id="rId1"/></sheets></workbook>'
    ),
    'xl/_rels/workbook.xml.rels': (
        '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>'
    ),
}

_MINI_DECK = {
    'ppt/presentation.xml': (
        '<?xml version="1.0"?><p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<p:sldIdLst><p:sldId id="256" r:id="rId2"/></p:sldIdLst></p:presentation>'
    ),
    'ppt/_rels/presentation.xml.rels': (
        '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId2" Target="slides/slide1.xml"/></Relationships>'
    ),
}


def _peak_bytes_of(call):
    import tracemalloc
    tracemalloc.start()
    try:
        with pytest.raises(HTTPException) as exc:
            call()
        return exc.value.status_code, tracemalloc.get_traced_memory()[1]
    finally:
        tracemalloc.stop()


@pytest.mark.parametrize('name,files,part', [
    ('bomb.xlsx', _MINI_WORKBOOK, 'xl/worksheets/sheet1.xml'),
    ('bomb.xlsx', _MINI_WORKBOOK, 'xl/sharedStrings.xml'),
    ('bomb.pptx', _MINI_DECK, 'ppt/slides/slide1.xml'),
], ids=['xlsx-sheet', 'xlsx-shared-strings', 'pptx-slide'])
def test_a_zip_bomb_is_refused_with_413_before_anything_is_inflated(name, files, part):
    payload = _zeros_zip({part: 200 * (1 << 20)}, files)
    assert len(payload) < 1_000_000  # it would pass the 1 MB attachment cap
    status, peak = _peak_bytes_of(lambda: extract_chat_document(request(name, payload)))
    assert status == 413
    assert peak < 16 * (1 << 20), f'peak allocation {peak} bytes'


def test_an_archive_with_too_many_members_is_refused_with_413():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w') as zf:
        for path, xml in _MINI_WORKBOOK.items():
            zf.writestr(path, xml)
        for i in range(2_001):
            zf.writestr(f'xl/junk/{i}.xml', '')
    with pytest.raises(HTTPException) as exc:
        extract_chat_document(request('many.xlsx', buffer.getvalue()))
    assert exc.value.status_code == 413


def test_many_parts_that_are_each_small_but_together_too_big_are_refused_with_413():
    parts = {f'xl/junk/{i}.bin': 9 * (1 << 20) for i in range(4)}  # 36 MB in total, none over the part cap
    status, peak = _peak_bytes_of(
        lambda: extract_chat_document(request('wide.xlsx', _zeros_zip(parts, _MINI_WORKBOOK)))
    )
    assert status == 413
    assert peak < 16 * (1 << 20)


def test_one_part_over_the_part_cap_is_refused_with_413():
    payload = _zeros_zip({'xl/worksheets/sheet1.xml': 25 * (1 << 20)}, _MINI_WORKBOOK)  # under the 30 MB total
    with pytest.raises(HTTPException) as exc:
        extract_chat_document(request('part.xlsx', payload))
    assert exc.value.status_code == 413


def test_an_entity_expansion_bomb_in_a_worksheet_is_refused_and_stays_bounded():
    # Python's default expat protections stay on: a billion-laughs payload ends
    # in a refusal instead of expanding without limit.
    entities = '<!ENTITY a0 "lol">' + ''.join(f'<!ENTITY a{i} "{"&a%d;" % (i - 1) * 10}">' for i in range(1, 12))
    sheet = (
        f'<?xml version="1.0"?><!DOCTYPE bomb [{entities}]>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
        '<row r="1"><c r="A1" t="inlineStr"><is><t>&a11;</t></is></c></row></sheetData></worksheet>'
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w') as zf:
        for path, xml in _MINI_WORKBOOK.items():
            zf.writestr(path, xml)
        zf.writestr('xl/worksheets/sheet1.xml', sheet)
    status, peak = _peak_bytes_of(lambda: extract_chat_document(request('laughs.xlsx', buffer.getvalue())))
    assert status == 422
    assert peak < 256 * (1 << 20), f'peak allocation {peak} bytes'


def test_an_absurd_column_reference_does_not_allocate_a_giant_row():
    sheet = (
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1">'
        '<c r="ZZZZZZZZZZ1" t="inlineStr"><is><t>far</t></is></c><c r="A1" t="inlineStr"><is><t>near</t></is></c>'
        '</row></sheetData></worksheet>'
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w') as zf:
        for path, xml in _MINI_WORKBOOK.items():
            zf.writestr(path, xml)
        zf.writestr('xl/worksheets/sheet1.xml', sheet)
    assert read('column.xlsx', buffer.getvalue()) == '[Ark]\nnear'


def test_a_workbook_that_takes_too_long_to_stream_is_refused_with_413(monkeypatch):
    from app import parsers
    monkeypatch.setattr(parsers, '_PARSE_DEADLINE_SECONDS', -1.0)  # already past
    with pytest.raises(HTTPException) as exc:
        extract_chat_document(request('slow.xlsx', _build_xlsx([_long_sheet('Data', 2000)])))
    assert exc.value.status_code == 413
    # The import pipeline passes no cap, so it has no deadline either.
    assert parsers._parse_local('xlsx', _build_xlsx([_long_sheet('Data', 2000)]), {}).endswith('Rad 01999\t12345')


def test_an_ordinary_workbook_is_inside_the_budget():
    from app.parsers import _zip_budget
    with zipfile.ZipFile(io.BytesIO(_build_xlsx([('Ark', [['a']])]))) as zf:
        _zip_budget(zf)  # does not raise


# The budget above bounds what a zip may declare. These bombs stay inside it:
# valid XML a few thousand times smaller on the wire than in memory, shaped to
# cost more than imports-api's 256 MB container once it is parsed.

def _repeated_element_zip(files: dict[str, str], part: str, head: str, unit: str, tail: str, size: int) -> bytes:
    buffer = io.BytesIO()
    chunk = (unit * ((1 << 20) // len(unit))).encode()
    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED, compresslevel=9) as zf:
        for path, xml in files.items():
            zf.writestr(path, xml)
        with zf.open(part, 'w', force_zip64=True) as out:
            out.write(head.encode())
            for _ in range(size // len(chunk)):
                out.write(chunk)
            out.write(tail.encode())
    return buffer.getvalue()


_SHEET_NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'


def test_a_part_that_is_read_whole_is_refused_before_it_can_fill_the_container():
    # 19 MB of empty elements is a 20 KB upload and costs several hundred MB as a tree.
    payload = _repeated_element_zip(
        _MINI_DECK, 'ppt/slides/slide1.xml',
        f'<p:sld xmlns:p="{_P_NS}"><p:cSld><p:spTree>', '<a/>', '</p:spTree></p:cSld></p:sld>', 19 * (1 << 20),
    )
    assert len(payload) < 1_000_000
    status, peak = _peak_bytes_of(lambda: extract_chat_document(request('tiny.pptx', payload)))
    assert status == 413
    assert peak < 16 * (1 << 20), f'peak allocation {peak} bytes'


def test_an_oversized_style_table_is_skipped_instead_of_refusing_the_workbook():
    import tracemalloc
    sheet = (
        f'<worksheet xmlns="{_SHEET_NS}"><sheetData><row r="1"><c r="A1" s="1"><v>45366</v></c></row></sheetData></worksheet>'
    )
    payload = _repeated_element_zip(
        {**_MINI_WORKBOOK, 'xl/worksheets/sheet1.xml': sheet}, 'xl/styles.xml',
        f'<styleSheet xmlns="{_SHEET_NS}">', '<a/>', '</styleSheet>', 19 * (1 << 20),
    )
    tracemalloc.start()
    try:
        text = read('styles.xlsx', payload)
        peak = tracemalloc.get_traced_memory()[1]
    finally:
        tracemalloc.stop()
    assert text == '[Ark]\n45366'  # no style table to read: the number stays a number
    assert peak < 16 * (1 << 20), f'peak allocation {peak} bytes'


@pytest.mark.parametrize('head,unit,tail', [
    (f'<worksheet xmlns="{_SHEET_NS}"><sheetData><row r="1">', '<c/>', '</row></sheetData></worksheet>'),
    (f'<worksheet xmlns="{_SHEET_NS}"><sheetData/><mergeCells>', '<mergeCell ref="A1:B1"/>', '</mergeCells></worksheet>'),
], ids=['one-row-of-millions-of-cells', 'millions-of-merged-ranges'])
def test_a_streamed_part_that_never_closes_what_it_opens_is_refused(head, unit, tail):
    payload = _repeated_element_zip(_MINI_WORKBOOK, 'xl/worksheets/sheet1.xml', head, unit, tail, 19 * (1 << 20))
    assert len(payload) < 1_000_000
    status, peak = _peak_bytes_of(lambda: extract_chat_document(request('wide.xlsx', payload)))
    assert status == 413
    assert peak < 64 * (1 << 20), f'peak allocation {peak} bytes'


def test_a_full_width_row_of_formula_cells_is_not_mistaken_for_a_bomb():
    # 16 384 cells of <c><f/><v/></c> keep 49 152 elements open until the row closes.
    from app.parsers import _col_index_to_letters
    cells = ''.join(f'<c r="{_col_index_to_letters(i)}1"><f>1</f><v>1</v></c>' for i in range(1, 16_385))
    sheet = f'<worksheet xmlns="{_SHEET_NS}"><sheetData><row r="1">{cells}</row></sheetData></worksheet>'
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w') as zf:
        for path, xml in {**_MINI_WORKBOOK, 'xl/worksheets/sheet1.xml': sheet}.items():
            zf.writestr(path, xml)
    text = read('wide.xlsx', buffer.getvalue())
    assert text.startswith('[Ark]\n' + '1\t' * 16_383 + '1\n')
    assert text.endswith('… 16284 flere formler utelatt')


def test_a_deck_that_points_every_slide_at_the_same_large_part_is_refused_on_time(monkeypatch):
    import time
    from app import parsers
    monkeypatch.setattr(parsers, '_PARSE_DEADLINE_SECONDS', 1.0)
    slides = 300
    relationships = 'xmlns="http://schemas.openxmlformats.org/package/2006/relationships"'
    sld_ids = ''.join(f'<p:sldId id="{256 + i}" r:id="rId{i + 2}"/>' for i in range(slides))
    deck_rels = ''.join(f'<Relationship Id="rId{i + 2}" Type="{_REL_TYPE}slide" Target="slides/slide{i}.xml"/>' for i in range(slides))
    parts = {
        'ppt/presentation.xml': (
            f'<?xml version="1.0"?><p:presentation xmlns:p="{_P_NS}" xmlns:r="{_R_NS}"><p:sldIdLst>{sld_ids}</p:sldIdLst></p:presentation>'
        ),
        'ppt/_rels/presentation.xml.rels': f'<?xml version="1.0"?><Relationships {relationships}>{deck_rels}</Relationships>',
    }
    slide = (
        f'<?xml version="1.0"?><p:sld xmlns:p="{_P_NS}" xmlns:a="{_A_NS}" xmlns:r="{_R_NS}">'
        f'<p:cSld><p:spTree>{_chart_frame("rId5")}</p:spTree></p:cSld></p:sld>'
    )
    slide_rels = (
        f'<?xml version="1.0"?><Relationships {relationships}>'
        f'<Relationship Id="rId5" Type="{_REL_TYPE}chart" Target="../charts/chart1.xml"/></Relationships>'
    )
    for i in range(slides):
        parts[f'ppt/slides/slide{i}.xml'] = slide
        parts[f'ppt/slides/_rels/slide{i}.xml.rels'] = slide_rels
    # One 1.5 MB chart (under the 2 MB limit for a part read whole) that every slide re-reads.
    payload = _repeated_element_zip(
        parts, 'ppt/charts/chart1.xml',
        f'<c:chartSpace xmlns:c="{_C_NS}"><c:chart><c:plotArea>', '<c:x/>', '</c:plotArea></c:chart></c:chartSpace>',
        3 * (1 << 19),
    )
    assert len(payload) < 1_000_000
    started = time.monotonic()
    with pytest.raises(HTTPException) as exc:
        extract_chat_document(request('loop.pptx', payload))
    assert exc.value.status_code == 413
    assert time.monotonic() - started < 15
    # The import pipeline passes no cap, so it has no deadline either.
    assert parsers._parse_local('pptx', _build_pptx([['Tittel']]), {}) == '[Slide 1]\nTittel'


# Strings that cost nothing to read once but are built again for every
# reference to them: a few KB on the wire, hundreds of MB in memory.

def _deflated(files: dict[str, str]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED, compresslevel=9) as zf:
        for path, xml in files.items():
            zf.writestr(path, xml)
    return buffer.getvalue()


def test_a_row_of_cells_that_all_name_one_long_shared_string_is_refused_before_it_is_joined():
    # 1 500 references to one 200 KB string are a few KB of upload and 300 MB once the row is joined.
    cells = '<c t="s"><v>0</v></c>' * 1_500
    payload = _deflated({
        **_MINI_WORKBOOK,
        'xl/sharedStrings.xml': f'<sst xmlns="{_SHEET_NS}"><si><t>{"A" * 200_000}</t></si></sst>',
        'xl/worksheets/sheet1.xml': f'<worksheet xmlns="{_SHEET_NS}"><sheetData><row r="1">{cells}</row></sheetData></worksheet>',
    })
    assert len(payload) < 10_000
    status, peak = _peak_bytes_of(lambda: extract_chat_document(request('refs.xlsx', payload)))
    assert status == 413
    assert peak < 16 * (1 << 20), f'peak allocation {peak} bytes'


def test_a_deck_that_names_one_long_slide_again_and_again_is_refused_before_the_text_piles_up():
    from app import parsers
    repeats = 100
    payload = _deflated({
        'ppt/presentation.xml': (
            f'<?xml version="1.0"?><p:presentation xmlns:p="{_P_NS}" xmlns:r="{_R_NS}"><p:sldIdLst>'
            + ''.join(f'<p:sldId id="{256 + i}" r:id="rId2"/>' for i in range(repeats))
            + '</p:sldIdLst></p:presentation>'
        ),
        'ppt/_rels/presentation.xml.rels': _MINI_DECK['ppt/_rels/presentation.xml.rels'],
        'ppt/slides/slide1.xml': (
            f'<?xml version="1.0"?><p:sld xmlns:p="{_P_NS}" xmlns:a="{_A_NS}"><p:cSld><p:spTree><p:sp><p:txBody><a:bodyPr/>'
            f'<a:p><a:r><a:t>{"A" * 300_000}</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:sld>'
        ),
    })
    assert len(payload) < 10_000
    status, peak = _peak_bytes_of(lambda: extract_chat_document(request('again.pptx', payload)))
    assert status == 413
    assert peak < 16 * (1 << 20), f'peak allocation {peak} bytes'
    # The import pipeline passes no cap and keeps every slide.
    assert len(parsers._parse_local('pptx', payload, {})) > repeats * 300_000


def test_a_very_long_sheet_name_is_cut_before_it_is_repeated_in_every_formula_line():
    import tracemalloc
    rows = ''.join(f'<row r="{i}"><c r="A{i}"><f>1+{i}</f><v>1</v></c></row>' for i in range(1, 101))
    payload = _deflated({
        'xl/workbook.xml': (
            f'<?xml version="1.0"?><workbook xmlns="{_SHEET_NS}" xmlns:r="{_R_NS}">'
            f'<sheets><sheet name="{"S" * 1_900_000}" sheetId="1" r:id="rId1"/></sheets></workbook>'
        ),
        'xl/_rels/workbook.xml.rels': _MINI_WORKBOOK['xl/_rels/workbook.xml.rels'],
        'xl/worksheets/sheet1.xml': f'<worksheet xmlns="{_SHEET_NS}"><sheetData>{rows}</sheetData></worksheet>',
    })
    tracemalloc.start()
    try:
        text = read('name.xlsx', payload)
        peak = tracemalloc.get_traced_memory()[1]
    finally:
        tracemalloc.stop()
    assert text.startswith('[' + 'S' * 255 + ']\n1\n')
    assert 'S' * 256 not in text
    assert peak < 16 * (1 << 20), f'peak allocation {peak} bytes'


# --- Files written by Excel and PowerPoint themselves -----------------------
# tests/fixtures/real/make-real-fixtures.ps1 authors these through COM; the
# values asserted here are the ones it writes. The hand-built zips above prove
# the code against what we think Office stores; these prove it against what
# Office stores.

REAL = Path(__file__).parent / 'fixtures' / 'real'

REAL_WORKBOOK_TEXT = '\n'.join([
    '[Budsjett]',
    'Budsjett 2024',
    'Frist\tBeløp\tPost\tAndel',
    '2024-03-15\t12000\tLeie\t34.3%',
    '2024-04-01\t4500\tStrøm\t12.9%',
    '2024-04-30\t9000\tLønn\t25.7%',
    '2024-05-15\t2500\tForsikring\t7.1%',
    '2024-06-01\t3000\tKontor\t8.6%',
    '2024-06-30\t1800\tProgramvare\t5.1%',
    '2024-07-15\t2200\tAnnet\t6.3%',
    'Sum\t35000\t\t100.0%',
    '',
    '[Sjekk]',
    'Over grense?\tTRUE',
    'Feil\t#DIV/0!',
    'Tid\t12:00',
    'Dato og tid\t2024-01-01 12:00',
    'Andel\t25.5%',
    'Støy\t0.3',
    '',
    '[Skjult] (skjult)',
    'Skjult notat',
    'Intern kode 7731',
    '',
    '[Formler]',
    'Budsjett!D3: =B3/$B$10',
    'Budsjett!D4: =B4/$B$10',
    'Budsjett!D5: =B5/$B$10',
    'Budsjett!D6: =B6/$B$10',
    'Budsjett!D7: =B7/$B$10',
    'Budsjett!D8: =B8/$B$10',
    'Budsjett!D9: =B9/$B$10',
    'Budsjett!B10: =SUM(B3:B9)',
    'Budsjett!D10: =SUM(D3:D9)',
    'Sjekk!B1: =Budsjett!B10>30000',
    'Sjekk!B2: =1/0',
    'Sjekk!B6: =0.1+0.2',
])

REAL_DECK_TEXT = '\n'.join([
    '[Slide 1]',
    'Kvartalsrapport Q1',
    'Intern gjennomgang',
    '',
    '[Slide 2]',
    'Resultater',
    'Region\tOmsetning\tVekst',
    'Nord\t120\t5 %',
    'Sør\t95\t3 %',
    '[Notater]',
    'Husk å nevne at Sør-regionen har ny leder fra mars.',
    '',
    '[Slide 3]',
    'Utvikling',
    '[Diagram: Omsetning per kvartal] Omsetning: 100, 120, 140; kategorier: Q1, Q2, Q3',
    '[Bilde: Firmalogo (test)]',
    '',
    '[Slide 4] (skjult)',
    'Reserveslide',
    'Ikke vis denne',
])


def test_a_workbook_written_by_excel_reads_the_way_excel_shows_it():
    result = extract_chat_document(request('regneark.xlsx', (REAL / 'regneark.xlsx').read_bytes()))['data']
    assert result['content'] == REAL_WORKBOOK_TEXT
    assert 'truncated' not in result


def test_the_excel_fixture_answers_the_three_live_questions_by_position():
    # "What date is in A3, what is the total in B10": row n is line n of the sheet.
    content = extract_chat_document(request('regneark.xlsx', (REAL / 'regneark.xlsx').read_bytes()))['data']['content']
    lines = content.split('\n')
    assert lines[3].split('\t')[0] == '2024-03-15'
    assert lines[10].split('\t')[1] == '35000'
    assert 'Budsjett!B10: =SUM(B3:B9)' in lines


def test_a_deck_written_by_powerpoint_reads_with_notes_table_chart_and_hidden_slide():
    result = extract_chat_document(request('presentasjon.pptx', (REAL / 'presentasjon.pptx').read_bytes()))['data']
    assert result['content'] == REAL_DECK_TEXT


def test_the_import_pipeline_still_reads_every_row_of_a_big_sheet():
    # The cut is a chat-turn limit; durable imports pass no cap and keep everything.
    from app.parsers import _parse_local
    metadata: dict = {}
    text = _parse_local('xlsx', _build_xlsx([_long_sheet('Data', 6000)]), metadata)
    assert 'avkortet' not in text
    assert text.endswith('Rad 05999\t12345')
    assert 'truncated' not in metadata
