import contextvars
import csv
import hashlib
import io
import json
import math
import posixpath
import re
import time
import zipfile
import xml.etree.ElementTree as ET
from datetime import date, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from typing import Any, NamedTuple

import httpx
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader

from app.config import get_settings
from app.schemas import ImportDocument


class UnsupportedFileTypeError(ValueError):
    pass


def _sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _parse_pdf(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    parts: list[str] = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return "\n".join(parts).strip()


def _parse_docx(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    # Preserve paragraph/table order: order quantities and policy revisions
    # frequently live in tables, not in document.paragraphs.
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    parts = []
    for element in document.element.body:
        if element.tag.endswith('}p'):
            parts.append(Paragraph(element, document).text)
        elif element.tag.endswith('}tbl'):
            table = Table(element, document)
            parts.extend('\t'.join(cell.text for cell in row.cells) for row in table.rows)
    return "\n".join(parts).strip()


def _parse_plain(content: bytes) -> str:
    return content.decode("utf-8", errors="ignore").strip()


def _parse_csv(content: bytes) -> str:
    decoded = content.decode("utf-8", errors="ignore")
    rows = list(csv.reader(io.StringIO(decoded)))
    return "\n".join("\t".join(row) for row in rows)


def _parse_json(content: bytes) -> tuple[str, dict[str, Any]]:
    payload = json.loads(content.decode("utf-8", errors="ignore"))
    if isinstance(payload, dict):
        keys = list(payload.keys())
    elif isinstance(payload, list):
        keys = ["<array>"]
    else:
        keys = [type(payload).__name__]
    return json.dumps(payload, ensure_ascii=False, indent=2), {"json_keys": keys}


def _parse_html(content: bytes) -> str:
    soup = BeautifulSoup(content, "lxml")
    return soup.get_text("\n", strip=True)


# ---------------------------------------------------------------------------
# XLSX / PPTX: both are zip containers of plain XML parts (OOXML), so the
# standard library's zipfile + xml.etree cover them without a new dependency
# (mirrors the Go extractor's archive/zip + encoding/xml approach).
#
# These are untrusted files. `_zip_budget` runs before any part is read,
# worksheets are streamed instead of built into a tree, and Python's default
# expat entity-expansion protections stay on.

_XLSX_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_OOXML_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_PPTX_MAIN_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
_DRAWINGML_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
_MARKUP_COMPAT_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"

_CELL_REF_RE = re.compile(r"^([A-Z]+)(\d+)$")

# Hard limits, checked on the zip directory before any part is inflated.
_ZIP_MAX_MEMBERS = 2000
_ZIP_MAX_PART_BYTES = 20 * 1024 * 1024
_ZIP_MAX_TOTAL_BYTES = 30 * 1024 * 1024

# A part that is parsed into a tree (every part except the streamed worksheets
# and shared strings) costs 20 to 35 times its size in memory, and imports-api
# runs in a 256 MB container: a 20 MB part of empty elements, which is a 20 KB
# upload, would be killed by the memory limit while the budget above lets it in.
_ZIP_MAX_TREE_PART_BYTES = 2 * 1024 * 1024
# The streamed parts free each row and string as soon as it is read, but one
# row, one string or one list of merged ranges is held whole until it closes.
# A real one has tens of thousands of elements at most (a row is 16 384 cells,
# three elements each when every cell carries a formula and a value).
_STREAM_MAX_OPEN_ELEMENTS = 100_000
# What a row, a deck or a name may build in memory out of references. A cell that
# points at a shared string costs nothing to read, but joining a row of 400 cells
# that all point at one 1 MB string builds 400 MB, and a deck whose slide list
# names one 2 MB slide again and again builds that text again and again. A chat
# turn takes 60 000 characters at most, so nothing near these limits is ever shown.
_XLSX_MAX_ROW_CHARS = 1_000_000
_PPTX_MAX_CHAT_CHARS = 200_000
_SHEET_NAME_MAX_CHARS = 255  # Excel allows 31; the name is repeated in every formula line

_XLSX_MAX_COLUMN = 16384  # XFD
_XLSX_MAX_ROW = 1_048_576
_XLSX_GAP_LINES = 50  # longer runs of missing rows become one jump marker
_XLSX_FORMULA_LINES = 100
_XLSX_FORMULA_TEXT_CHARS = 300
_XLSX_FORMULA_SECTION_CHARS = 3_500  # when the caller caps the text (chat)
_XLSX_MAX_OMITTED_SHEET_MARKERS = 10

_CHART_MAX_SERIES = 12
_CHART_MAX_POINTS = 100

# Wall-clock budget for streaming a workbook, so a worst-case sheet ends in a
# refusal well before the gateway's own 25 s timeout.
_PARSE_DEADLINE_SECONDS = 15.0

# The same budget for the parts that are read whole (a deck is read part by
# part, and a crafted one can point every slide at the same large part).
_READ_DEADLINE: contextvars.ContextVar[float | None] = contextvars.ContextVar(
    "office_read_deadline", default=None,
)

# Clark-notation tags, built once: the cell loop runs hundreds of thousands of times.
_X_SST = f"{{{_XLSX_MAIN_NS}}}sst"
_X_SI = f"{{{_XLSX_MAIN_NS}}}si"
_X_T = f"{{{_XLSX_MAIN_NS}}}t"
_X_IS = f"{{{_XLSX_MAIN_NS}}}is"
_X_SHEET_DATA = f"{{{_XLSX_MAIN_NS}}}sheetData"
_X_ROW = f"{{{_XLSX_MAIN_NS}}}row"
_X_C = f"{{{_XLSX_MAIN_NS}}}c"
_X_F = f"{{{_XLSX_MAIN_NS}}}f"
_X_V = f"{{{_XLSX_MAIN_NS}}}v"


class _MalformedOfficeDocument(ValueError):
    """A .xlsx/.pptx zip is missing or has malformed OOXML parts."""


class OfficeArchiveTooLargeError(ValueError):
    """A .xlsx/.pptx zip declares more content than we are willing to inflate."""


def _zip_budget(zf: zipfile.ZipFile) -> None:
    """Refuse an archive from its directory alone, before any part is read.

    `ZipInfo.file_size` is the declared uncompressed size, and `zipfile` never
    hands out more than that, so checking it first bounds every later read.
    A zip bomb of a few hundred KB ends here without allocating anything."""
    infos = zf.infolist()
    if len(infos) > _ZIP_MAX_MEMBERS:
        raise OfficeArchiveTooLargeError("The archive has too many parts")
    total = 0
    for info in infos:
        if info.file_size > _ZIP_MAX_PART_BYTES:
            raise OfficeArchiveTooLargeError("A part of the archive is too large")
        total += info.file_size
        if total > _ZIP_MAX_TOTAL_BYTES:
            raise OfficeArchiveTooLargeError("The archive expands to too much content")


def _has_part(zf: zipfile.ZipFile, path: str) -> bool:
    try:
        zf.getinfo(path)
    except KeyError:
        return False
    return True


def _check_read_deadline() -> None:
    deadline = _READ_DEADLINE.get()
    if deadline is not None and time.monotonic() > deadline:
        raise OfficeArchiveTooLargeError("The document took too long to read")


def _read_zip_xml(zf: zipfile.ZipFile, path: str) -> ET.Element | None:
    _check_read_deadline()
    try:
        info = zf.getinfo(path)
    except KeyError:
        return None
    if info.file_size > _ZIP_MAX_TREE_PART_BYTES:
        raise OfficeArchiveTooLargeError("A part of the archive is too large to read safely")
    try:
        return ET.fromstring(zf.read(info))
    except ET.ParseError as exc:
        raise _MalformedOfficeDocument(f"Malformed XML part: {path}") from exc


def _iter_elements(
    zf: zipfile.ZipFile,
    path: str,
    tag: str,
    container_tag: str,
    deadline: float | None = None,
):
    """Stream the `tag` children of `container_tag` out of a large part.

    Each element is cleared once the caller is done with it, and so is its
    container, so memory stays flat however many rows or strings the part
    holds. Past `deadline` (a `time.monotonic()` value) the stream is refused,
    and so is a part that keeps opening elements without ever closing a row or
    string (one `<row>` of millions of cells is held whole until it ends)."""
    if not _has_part(zf, path):
        return
    if deadline is not None and time.monotonic() > deadline:
        raise OfficeArchiveTooLargeError("The document took too long to read")
    container: ET.Element | None = None
    events = 0
    depth = 0
    pending = 0  # elements opened since the last point where memory was released
    try:
        with zf.open(path) as stream:
            for event, elem in ET.iterparse(stream, events=("start", "end")):
                events += 1
                if deadline is not None and not events & 0xFFF and time.monotonic() > deadline:
                    raise OfficeArchiveTooLargeError("The document took too long to read")
                if event == "start":
                    depth += 1
                    pending += 1
                    if pending > _STREAM_MAX_OPEN_ELEMENTS:
                        raise OfficeArchiveTooLargeError("A row or string of the document holds too much content")
                    if container is None and elem.tag == container_tag:
                        container = elem
                    continue
                depth -= 1
                if elem.tag == tag:
                    yield elem
                    elem.clear()
                    if container is not None:
                        container.clear()
                    pending = 0
                elif depth == 1:
                    # A top-level part we do not read (merged ranges, hyperlinks,
                    # a column list): free it as soon as it closes.
                    elem.clear()
                    pending = 0
    except ET.ParseError as exc:
        raise _MalformedOfficeDocument(f"Malformed XML part: {path}") from exc


def _relationship_entries(zf: zipfile.ZipFile, rels_path: str) -> dict[str, tuple[str, str, str]]:
    """Map relationship id -> (type, target, target mode), from a `_rels/*.rels` part."""
    root = _read_zip_xml(zf, rels_path)
    if root is None:
        return {}
    return {
        rel.get("Id", ""): (rel.get("Type", ""), rel.get("Target", ""), rel.get("TargetMode", ""))
        for rel in root.findall(f"{{{_PACKAGE_REL_NS}}}Relationship")
    }


def _relationship_targets(zf: zipfile.ZipFile, rels_path: str) -> dict[str, str]:
    """Map relationship id -> target path, from a `_rels/*.rels` part."""
    return {
        rid: target for rid, (_type, target, _mode) in _relationship_entries(zf, rels_path).items()
    }


def _rels_path_for(part_path: str) -> str:
    return posixpath.join(posixpath.dirname(part_path), "_rels", posixpath.basename(part_path) + ".rels")


def _resolve_part_path(base_part: str, target: str) -> str:
    """Zip path a relationship `target` names, relative to the part that holds it."""
    if target.startswith("/"):
        return posixpath.normpath(target.lstrip("/"))
    return posixpath.normpath(posixpath.join(posixpath.dirname(base_part), target))


def _col_letters_to_index(letters: str) -> int:
    """Spreadsheet column letters ("A", "Z", "AA", ...) to a 1-based index."""
    index = 0
    for char in letters:
        index = index * 26 + (ord(char) - ord("A") + 1)
    return index


def _col_index_to_letters(index: int) -> str:
    letters = ""
    while index > 0:
        index, remainder = divmod(index - 1, 26)
        letters = chr(ord("A") + remainder) + letters
    return letters


def _xlsx_shared_strings(zf: zipfile.ZipFile, deadline: float | None = None) -> list[str]:
    return [
        "".join(t.text or "" for t in si.iter(_X_T))
        for si in _iter_elements(zf, "xl/sharedStrings.xml", _X_SI, _X_SST, deadline)
    ]


# --- Number formats ---------------------------------------------------------
#
# A cell stores a bare number; whether it is a date, a time or a percentage
# lives only in the style (`s`) the cell points at. Read as plain numbers, a
# date is a five-digit integer and 30 % is 0.3.

class _NumberFormat(NamedTuple):
    kind: str  # "general" | "date" | "datetime" | "time" | "percent"
    decimals: int = 0  # percent only
    elapsed: bool = False  # [h]:mm style time, which may pass 24 hours


_GENERAL_FORMAT = _NumberFormat("general")

_BUILTIN_NUMBER_FORMATS: dict[int, _NumberFormat] = {
    9: _NumberFormat("percent", 0),
    10: _NumberFormat("percent", 2),
    22: _NumberFormat("datetime"),
    46: _NumberFormat("time", elapsed=True),
}
for _format_id in (14, 15, 16, 17, 27, 28, 29, 30, 31, 36, *range(50, 59)):
    _BUILTIN_NUMBER_FORMATS[_format_id] = _NumberFormat("date")
for _format_id in (18, 19, 20, 21, 32, 33, 34, 35, 45, 47):
    _BUILTIN_NUMBER_FORMATS[_format_id] = _NumberFormat("time")

_QUOTED_LITERAL_RE = re.compile(r'"[^"]*"')
_ESCAPED_CHAR_RE = re.compile(r"\\.")
_PADDING_RE = re.compile(r"[_*].")
_BRACKET_RE = re.compile(r"\[([^\]]*)\]")
_ELAPSED_RE = re.compile(r"^(h+|m+|s+)$", re.IGNORECASE)
_AM_PM_RE = re.compile(r"am/pm|a/p", re.IGNORECASE)
_GENERAL_WORD_RE = re.compile(r"general", re.IGNORECASE)
_DATE_TIME_TOKEN_RE = re.compile(r"y+|m+|d+|h+|s+")


def _classify_number_format(code: str) -> _NumberFormat:
    """Date, time, percent or general, from a custom number-format code."""
    elapsed = False

    def bracket(match: re.Match) -> str:
        nonlocal elapsed
        inner = match.group(1)
        if _ELAPSED_RE.match(inner):
            elapsed = True
            return inner  # [h] still means hours
        return ""  # [Red], [$-409], [>=100] and friends are not part of the shape

    cleaned = _QUOTED_LITERAL_RE.sub("", code)
    cleaned = _ESCAPED_CHAR_RE.sub("", cleaned)
    cleaned = _PADDING_RE.sub("", cleaned)
    cleaned = _BRACKET_RE.sub(bracket, cleaned)
    cleaned = _GENERAL_WORD_RE.sub("", cleaned)
    section = next((part for part in cleaned.split(";") if part.strip()), "")

    if "%" in section:
        head = section.split("%")[0]
        decimals = 0
        if "." in head:
            decimals = len(re.match(r"[0#?]*", head.split(".", 1)[1]).group(0))
        return _NumberFormat("percent", decimals)

    tokens = _DATE_TIME_TOKEN_RE.findall(_AM_PM_RE.sub("", section).lower())
    has_date = has_time = False
    for index, token in enumerate(tokens):
        letter = token[0]
        if letter in "yd":
            has_date = True
        elif letter in "hs":
            has_time = True
        else:  # "m" is minutes next to hours or seconds, months everywhere else
            before = tokens[index - 1][0] if index else ""
            after = tokens[index + 1][0] if index + 1 < len(tokens) else ""
            if before == "h" or after == "s":
                has_time = True
            else:
                has_date = True
    if has_date and has_time:
        return _NumberFormat("datetime")
    if has_date:
        return _NumberFormat("date")
    if has_time:
        return _NumberFormat("time", elapsed=elapsed)
    return _GENERAL_FORMAT


def _xlsx_styles(zf: zipfile.ZipFile) -> list[_NumberFormat]:
    """The number format of each cell style (`cellXfs`), by style index."""
    try:
        if zf.getinfo("xl/styles.xml").file_size > _ZIP_MAX_TREE_PART_BYTES:
            # A bloated style table is not worth refusing the workbook for: its
            # cells print as plain numbers, the way they did before styles were read.
            return []
    except KeyError:
        return []
    root = _read_zip_xml(zf, "xl/styles.xml")
    if root is None:
        return []
    ns = {"m": _XLSX_MAIN_NS}
    custom: dict[int, str] = {}
    for fmt in root.findall("m:numFmts/m:numFmt", ns):
        try:
            custom[int(fmt.get("numFmtId", ""))] = fmt.get("formatCode", "")
        except ValueError:
            continue
    styles: list[_NumberFormat] = []
    for xf in root.findall("m:cellXfs/m:xf", ns):
        try:
            format_id = int(xf.get("numFmtId", "0"))
        except ValueError:
            format_id = 0
        if format_id in custom:
            styles.append(_classify_number_format(custom[format_id]))
        else:
            styles.append(_BUILTIN_NUMBER_FORMATS.get(format_id, _GENERAL_FORMAT))
    return styles


def _xlsx_uses_1904_dates(zf: zipfile.ZipFile) -> bool:
    workbook = _read_zip_xml(zf, "xl/workbook.xml")
    properties = workbook.find(f"{{{_XLSX_MAIN_NS}}}workbookPr") if workbook is not None else None
    return properties is not None and properties.get("date1904", "").lower() in {"1", "true"}


def _excel_serial_to_text(
    serial: float,
    kind: str,
    date1904: bool = False,
    elapsed: bool = False,
) -> str | None:
    """A spreadsheet serial as ISO text (YYYY-MM-DD, with HH:MM[:SS] for
    datetimes and times), or None when it is not a representable date."""
    if not math.isfinite(serial) or serial < 0 or serial >= 2_958_466:
        return None
    days = math.floor(serial)
    seconds = round((serial - days) * 86400)
    if seconds >= 86400:
        days, seconds = days + 1, seconds - 86400
    hours, remainder = divmod(seconds, 3600)
    minutes, secs = divmod(remainder, 60)
    if kind == "time" and elapsed:
        hours, remainder = divmod(round(serial * 86400), 3600)
        minutes, secs = divmod(remainder, 60)
    clock = f"{hours:02d}:{minutes:02d}" + (f":{secs:02d}" if secs else "")
    if kind == "time":
        return clock
    try:
        if date1904:
            day_text = (date(1904, 1, 1) + timedelta(days=days)).isoformat()
        elif days == 0:
            day_text = "1900-01-00"  # what Excel shows for a zero date
        elif days < 60:
            day_text = (date(1899, 12, 31) + timedelta(days=days)).isoformat()
        elif days == 60:
            day_text = "1900-02-29"  # Excel's leap day that never existed
        else:
            day_text = (date(1899, 12, 30) + timedelta(days=days)).isoformat()
    except OverflowError:
        return None
    return f"{day_text} {clock}" if kind == "datetime" else day_text


_INTEGER_TEXT_RE = re.compile(r"^-?\d+$")


def _normalise_number(text: str) -> str:
    """Drop binary-float noise (0.30000000000000004 -> 0.3): Excel keeps 15
    significant digits, so a longer tail is never information."""
    stripped = text.strip()
    if len(stripped) <= 15 and "e" not in stripped and "E" not in stripped:
        return stripped  # too short to carry a noisy tail
    if _INTEGER_TEXT_RE.match(stripped):
        return stripped
    try:
        value = float(stripped)
    except ValueError:
        return text
    if not math.isfinite(value):
        return text
    shown = format(value, ".15g")
    if "e" in shown:
        try:
            exact = Decimal(shown)
        except InvalidOperation:
            return shown
        if abs(exact.adjusted()) <= 20:
            return format(exact, "f")
    return shown


def _percent_text(raw: str, decimals: int) -> str | None:
    try:
        shown = Decimal(raw.strip()) * 100
        quantum = Decimal(1).scaleb(-decimals)
        return format(shown.quantize(quantum, rounding=ROUND_HALF_UP), "f") + "%"
    except (InvalidOperation, ValueError):
        return None


def _format_number(raw: str, number_format: _NumberFormat | None, date1904: bool) -> str:
    if number_format is not None and number_format.kind != "general":
        if number_format.kind == "percent":
            text = _percent_text(raw, number_format.decimals)
        else:
            try:
                number = float(raw)
            except ValueError:
                return raw
            text = _excel_serial_to_text(number, number_format.kind, date1904, number_format.elapsed)
        if text is not None:
            return text
    return _normalise_number(raw)


# --- Worksheets ---------------------------------------------------------------

_SHEET_STATES_HIDDEN = {"hidden", "veryHidden"}


def _xlsx_sheet_parts(zf: zipfile.ZipFile) -> list[tuple[str, str, str]]:
    """Return `[(sheet_name, worksheet_zip_path, state), ...]` in workbook order."""
    workbook = _read_zip_xml(zf, "xl/workbook.xml")
    if workbook is None:
        raise _MalformedOfficeDocument("xl/workbook.xml is missing")
    targets = _relationship_targets(zf, "xl/_rels/workbook.xml.rels")
    ns = {"m": _XLSX_MAIN_NS}
    parts: list[tuple[str, str, str]] = []
    for sheet in workbook.findall(".//m:sheets/m:sheet", ns):
        name = sheet.get("name", "")[:_SHEET_NAME_MAX_CHARS]
        rid = sheet.get(f"{{{_OOXML_REL_NS}}}id", "")
        target = targets.get(rid, "")
        if not target:
            continue
        # Target is normally relative to `xl/` ("worksheets/sheet1.xml"), but
        # OPC also allows a package-root-absolute form ("/xl/worksheets/...").
        # Strip the leading slash *before* checking for the "xl/" prefix, or
        # an absolute target doubles up into "xl/xl/worksheets/..." and the
        # sheet silently reads as empty instead of being found.
        target = target.lstrip('/')
        path = target if target.startswith("xl/") else f"xl/{target}"
        parts.append((name, path, sheet.get("state", "visible")))
    return parts


def _xlsx_formula_text(formula: ET.Element) -> str:
    body = " ".join((formula.text or "").split())
    if len(body) > _XLSX_FORMULA_TEXT_CHARS:
        body = body[:_XLSX_FORMULA_TEXT_CHARS] + "…"
    if body:
        return "=" + body
    # A shared-formula child stores no text of its own: the master cell holds it.
    return "=(delt formel)" if formula.get("t") == "shared" else "=(formel)"


def _xlsx_cell_text(
    cell: ET.Element,
    shared: list[str],
    styles: list[_NumberFormat] | None = None,
    date1904: bool = False,
) -> str:
    cell_type = cell.get("t")
    if cell_type == "inlineStr":
        inline = cell.find(_X_IS)
        return "".join(t.text or "" for t in inline.iter(_X_T)) if inline is not None else ""
    value = cell.find(_X_V)
    raw = value.text if value is not None else None
    formula = cell.find(_X_F)
    if formula is not None and not raw and not (cell_type == "str" and value is not None):
        # A formula nobody calculated (openpyxl writes `<f>..</f><v></v>`):
        # the formula is the only content, so show it instead of a blank.
        return _xlsx_formula_text(formula)
    if raw is None:
        return ""
    if cell_type == "s":
        try:
            return shared[int(raw)]
        except (ValueError, IndexError):
            return ""
    if cell_type == "b":
        flag = raw.strip()
        return "TRUE" if flag == "1" else "FALSE" if flag == "0" else raw
    if cell_type in {"str", "e", "d"}:
        return raw
    number_format = None
    if styles:
        try:
            style_index = int(cell.get("s", "0"))
        except ValueError:
            style_index = 0
        if 0 <= style_index < len(styles):
            number_format = styles[style_index]
    return _format_number(raw, number_format, date1904)


class _FormulaList:
    """The formulas behind the numbers, for the `[Formler]` section."""

    def __init__(self) -> None:
        self.lines: list[str] = []
        self.total = 0

    def add(self, sheet_name: str, ref: str, formula: ET.Element) -> None:
        if not (formula.text or "").strip():
            return  # a shared-formula child: its master is listed instead
        self.total += 1
        if len(self.lines) >= _XLSX_FORMULA_LINES:
            return
        if re.fullmatch(r"[^\W\d][\w.]*", sheet_name):
            qualified = sheet_name
        else:
            qualified = "'" + sheet_name.replace("'", "''") + "'"
        line = f"{qualified}!{ref}: {_xlsx_formula_text(formula)}"
        covered = formula.get("ref", "")
        if ":" in covered:
            kind = "delt formel" if formula.get("t") == "shared" else "matriseformel"
            line += f" ({kind} for {covered})"
        self.lines.append(line)

    def render(self, max_chars: int | None) -> str:
        if not self.total:
            return ""
        shown: list[str] = []
        used = len("[Formler]")
        for line in self.lines:
            if max_chars is not None and used + len(line) + 1 > max_chars:
                break
            shown.append(line)
            used += len(line) + 1
        if self.total > len(shown):
            shown.append(f"… {self.total - len(shown)} flere formler utelatt")
        return "\n".join(["[Formler]", *shown])


class _RowScan:
    """Lets the consumer of a row stream say it no longer needs the row text
    (the chat limit is reached): the stream then only reports whether each row
    holds anything, and keeps collecting formulas."""

    light = False


def _xlsx_cell_has_content(cell: ET.Element) -> bool:
    value = cell.find(_X_V)
    return bool(value is not None and value.text) or cell.find(_X_IS) is not None or cell.find(_X_F) is not None


def _xlsx_iter_rows(
    zf: zipfile.ZipFile,
    path: str,
    shared: list[str],
    styles: list[_NumberFormat],
    date1904: bool,
    sheet_name: str,
    formulas: _FormulaList,
    scan: _RowScan,
    deadline: float | None = None,
):
    """Yield `(row_number, tab-separated text)` per worksheet row, streaming."""
    last_row = 0
    for row in _iter_elements(zf, path, _X_ROW, _X_SHEET_DATA, deadline):
        try:
            row_number = int(row.get("r", ""))
        except ValueError:
            row_number = last_row + 1
        if not 0 < row_number <= _XLSX_MAX_ROW:
            row_number = last_row + 1
        last_row = max(last_row, row_number)
        light = scan.light
        by_column: dict[int, str] = {}
        row_chars = 0
        has_content = False
        next_column = 1
        for cell in row:
            if cell.tag != _X_C:
                continue
            ref = cell.get("r", "")
            match = _CELL_REF_RE.match(ref)
            column = _col_letters_to_index(match.group(1)) if match else next_column
            if not 0 < column <= _XLSX_MAX_COLUMN:
                column = next_column
            next_column = column + 1
            formula = cell.find(_X_F)
            if formula is not None:
                formulas.add(sheet_name, ref or f"{_col_index_to_letters(column)}{row_number}", formula)
            if light:
                has_content = has_content or _xlsx_cell_has_content(cell)
            else:
                cell_text = _xlsx_cell_text(cell, shared, styles, date1904)
                row_chars += len(cell_text) + 1
                if row_chars > _XLSX_MAX_ROW_CHARS:
                    raise OfficeArchiveTooLargeError("A row of the document holds too much text")
                by_column[column] = cell_text
        if light:
            yield row_number, "…" if has_content else ""
            continue
        last_column = max(by_column, default=0)
        # Formatted-but-empty cells would otherwise leave a run of tabs.
        yield row_number, "\t".join(by_column.get(i, "") for i in range(1, last_column + 1)).rstrip("\t")


def _parse_xlsx(
    content: bytes,
    *,
    max_chars: int | None = None,
    info: dict[str, Any] | None = None,
) -> str:
    """Sheet names and cell values as tab-separated rows per sheet, in
    workbook order. Numbers come back the way Excel shows them (dates as ISO
    dates, percentages with a % sign), formulas nobody calculated come back as
    `=SUM(B2:B9)`, and every formula is listed once more after the grid.

    Row n is line n of its sheet (missing and empty rows keep their place), so
    a question about "B10" has one answer. With `max_chars` the grid is cut on
    a row boundary once it would pass that many characters, with a marker, and
    `info["truncated"]` is set."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            _zip_budget(zf)
            # Only a chat turn is latency-bound; a durable import may take as long as it takes.
            deadline = time.monotonic() + _PARSE_DEADLINE_SECONDS if max_chars is not None else None
            shared = _xlsx_shared_strings(zf, deadline)
            styles = _xlsx_styles(zf)
            date1904 = _xlsx_uses_1904_dates(zf)
            formulas = _FormulaList()
            scan = _RowScan()
            sections: list[str] = []
            budget_left = max_chars
            exhausted = False
            truncated = False
            omitted_sheet_markers = 0
            omitted_sheets = 0
            for name, path, state in _xlsx_sheet_parts(zf):
                header = f"[{name}]" + (" (skjult)" if state in _SHEET_STATES_HIDDEN else "")
                lines: list[str] = []
                if not exhausted:
                    lines.append(header)
                    if budget_left is not None:
                        budget_left -= len(header) + 1
                omitted_rows = 0
                last_row = 0
                skipped_rows = 0
                rows = _xlsx_iter_rows(zf, path, shared, styles, date1904, name, formulas, scan, deadline)
                for row_number, text in rows:
                    skipped_rows += max(row_number - last_row - 1, 0)  # rows the file never wrote
                    last_row = max(last_row, row_number)
                    if not text:
                        skipped_rows += 1
                        continue
                    if exhausted:
                        omitted_rows += 1
                        continue
                    if skipped_rows > _XLSX_GAP_LINES:
                        chunk = [f"[hopper til rad {row_number}]", text]
                    else:
                        chunk = [""] * skipped_rows + [text]
                    skipped_rows = 0
                    cost = sum(len(line) + 1 for line in chunk)
                    if budget_left is not None and cost > budget_left:
                        exhausted = True
                        scan.light = True  # only counting from here on
                        omitted_rows += 1
                        continue
                    lines.extend(chunk)
                    if budget_left is not None:
                        budget_left -= cost
                if omitted_rows:
                    truncated = True
                    marker = f"[… avkortet: {omitted_rows} rader utelatt i «{name}»]"
                    if lines:
                        lines.append(marker)
                    elif omitted_sheet_markers < _XLSX_MAX_OMITTED_SHEET_MARKERS:
                        omitted_sheet_markers += 1
                        lines.append(marker)
                    else:
                        omitted_sheets += 1
                if lines:
                    sections.append("\n".join(lines))
            if omitted_sheets:
                sections.append(f"[… avkortet: {omitted_sheets} flere ark utelatt]")
            rendered = formulas.render(_XLSX_FORMULA_SECTION_CHARS if max_chars is not None else None)
            if rendered:
                sections.append(rendered)
            if truncated and info is not None:
                info["truncated"] = True
            return "\n\n".join(sections).strip()
    except zipfile.BadZipFile as exc:
        raise _MalformedOfficeDocument("Not a valid .xlsx archive") from exc


# --- Presentations ------------------------------------------------------------

def _pptx_slide_parts(zf: zipfile.ZipFile) -> list[str]:
    """Return slide XML zip paths in the presentation's own slide order."""
    presentation = _read_zip_xml(zf, "ppt/presentation.xml")
    if presentation is None:
        raise _MalformedOfficeDocument("ppt/presentation.xml is missing")
    targets = _relationship_targets(zf, "ppt/_rels/presentation.xml.rels")
    ns = {"p": _PPTX_MAIN_NS}
    paths: list[str] = []
    for sld_id in presentation.findall(".//p:sldIdLst/p:sldId", ns):
        rid = sld_id.get(f"{{{_OOXML_REL_NS}}}id", "")
        target = targets.get(rid, "")
        if not target:
            continue
        # See the matching comment in _xlsx_sheet_parts: strip before prefixing
        # so a package-root-absolute target doesn't double up to "ppt/ppt/...".
        target = target.lstrip('/')
        paths.append(target if target.startswith("ppt/") else f"ppt/{target}")
    return paths


def _pptx_paragraph_text(paragraph: ET.Element) -> str:
    parts: list[str] = []
    for child in paragraph:
        if child.tag in (f"{{{_DRAWINGML_NS}}}r", f"{{{_DRAWINGML_NS}}}fld"):
            parts.append("".join(t.text or "" for t in child.findall(f"{{{_DRAWINGML_NS}}}t")))
        elif child.tag == f"{{{_DRAWINGML_NS}}}br":
            parts.append("\n")
    return "".join(parts)


def _pptx_paragraph_lines(body: ET.Element | None) -> list[str]:
    if body is None:
        return []
    texts = (_pptx_paragraph_text(p) for p in body.findall(f"{{{_DRAWINGML_NS}}}p"))
    return [text for text in texts if text.strip()]


def _chart_number(raw: str, format_code: str, date1904: bool) -> str:
    number_format = _classify_number_format(format_code) if format_code else _GENERAL_FORMAT
    return _format_number(raw, number_format, date1904)


def _chart_points(container: ET.Element | None, date1904: bool) -> list[str]:
    """The cached points of a series' values or categories, in index order."""
    if container is None:
        return []
    ns = f"{{{_CHART_NS}}}"
    cache = next(
        (e for e in container.iter() if e.tag in (
            f"{ns}numCache", f"{ns}strCache", f"{ns}numLit", f"{ns}strLit", f"{ns}multiLvlStrCache",
        )),
        None,
    )
    if cache is not None and cache.tag == f"{ns}multiLvlStrCache":
        cache = cache.find(f"{ns}lvl")  # the innermost level is the first one written
    if cache is None:
        return []
    format_node = cache.find(f"{ns}formatCode")
    format_code = (format_node.text or "") if format_node is not None else ""
    is_numeric = cache.tag in (f"{ns}numCache", f"{ns}numLit")
    points: dict[int, str] = {}
    for point in cache.findall(f"{ns}pt"):
        try:
            index = int(point.get("idx", ""))
        except ValueError:
            continue
        v = point.find(f"{ns}v")
        raw = (v.text or "") if v is not None else ""
        points[index] = _chart_number(raw, format_code, date1904) if is_numeric and raw else raw
    count_node = cache.find(f"{ns}ptCount")
    try:
        declared = int(count_node.get("val", "0")) if count_node is not None else 0
    except ValueError:
        declared = 0
    size = max(declared, max(points, default=-1) + 1)
    values = [points.get(i, "") for i in range(min(size, _CHART_MAX_POINTS))]
    if size > _CHART_MAX_POINTS:
        values.append(f"… (+{size - _CHART_MAX_POINTS})")
    return values


def _pptx_chart_text(chart: ET.Element) -> str:
    """`[Diagram: title] series: v1, v2; ...; kategorier: c1, c2` from the
    values PowerPoint caches inside the chart part."""
    ns = f"{{{_CHART_NS}}}"
    date1904_node = chart.find(f"{ns}date1904")
    date1904 = date1904_node is not None and date1904_node.get("val", "1").lower() in {"1", "true"}
    title_node = chart.find(f"{ns}chart/{ns}title")
    title = ""
    if title_node is not None:
        title = " ".join("".join(t.text or "" for t in title_node.iter(f"{{{_DRAWINGML_NS}}}t")).split())
    parts: list[str] = []
    categories: list[str] = []
    series = list(chart.iter(f"{ns}ser"))
    for index, ser in enumerate(series[:_CHART_MAX_SERIES], start=1):
        name_node = ser.find(f"{ns}tx")
        name = ""
        if name_node is not None:
            name_value = next(name_node.iter(f"{ns}v"), None)
            name = " ".join((name_value.text or "").split()) if name_value is not None else ""
        values_node = ser.find(f"{ns}val")
        if values_node is None:
            values_node = ser.find(f"{ns}yVal")
        category_node = ser.find(f"{ns}cat")
        if category_node is None:
            category_node = ser.find(f"{ns}xVal")
        if not categories:
            categories = _chart_points(category_node, date1904)
        parts.append(f"{name or f'Serie {index}'}: {', '.join(_chart_points(values_node, date1904))}")
    if len(series) > _CHART_MAX_SERIES:
        parts.append(f"… (+{len(series) - _CHART_MAX_SERIES} serier)")
    if categories:
        parts.append(f"kategorier: {', '.join(categories)}")
    head = f"[Diagram: {title}]" if title else "[Diagram]"
    return f"{head} {'; '.join(parts)}" if parts else head


def _pptx_table_lines(table: ET.Element) -> list[str]:
    ns = f"{{{_DRAWINGML_NS}}}"
    lines: list[str] = []
    for row in table.findall(f"{ns}tr"):
        cells = []
        for cell in row.findall(f"{ns}tc"):
            paragraphs = _pptx_paragraph_lines(cell.find(f"{ns}txBody"))
            cells.append(" ".join(" ".join(paragraphs).split()))
        text = "\t".join(cells).rstrip("\t")
        if text.strip():
            lines.append(text)
    return lines


def _pptx_walk_shapes(
    zf: zipfile.ZipFile,
    slide_path: str,
    rels: dict[str, tuple[str, str, str]],
    node: ET.Element,
    title: list[str],
    body: list[str],
) -> None:
    """Collect a slide's text in the order the shapes sit in the shape tree."""
    p = f"{{{_PPTX_MAIN_NS}}}"
    a = f"{{{_DRAWINGML_NS}}}"
    for child in node:
        if child.tag == f"{p}sp":
            placeholder = child.find(f"{p}nvSpPr/{p}nvPr/{p}ph")
            kind = placeholder.get("type", "") if placeholder is not None else ""
            if kind == "sldNum":
                continue  # the slide number field is not content
            lines = _pptx_paragraph_lines(child.find(f"{p}txBody"))
            (title if kind in {"title", "ctrTitle"} else body).extend(lines)
        elif child.tag == f"{p}pic":
            properties = child.find(f"{p}nvPicPr/{p}cNvPr")
            alt = " ".join((properties.get("descr", "") if properties is not None else "").split())
            if alt:
                body.append(f"[Bilde: {alt[:300]}]")
        elif child.tag == f"{p}graphicFrame":
            data = child.find(f"{a}graphic/{a}graphicData")
            if data is None:
                continue
            table = data.find(f"{a}tbl")
            if table is not None:
                body.extend(_pptx_table_lines(table))
            chart_ref = data.find(f"{{{_CHART_NS}}}chart")
            if chart_ref is not None:
                relationship = rels.get(chart_ref.get(f"{{{_OOXML_REL_NS}}}id", ""))
                if relationship and relationship[0].endswith("/chart") and relationship[2] != "External":
                    chart = _read_zip_xml(zf, _resolve_part_path(slide_path, relationship[1]))
                    if chart is not None:
                        body.append(_pptx_chart_text(chart))
        elif child.tag == f"{p}grpSp":
            _pptx_walk_shapes(zf, slide_path, rels, child, title, body)
        elif child.tag == f"{{{_MARKUP_COMPAT_NS}}}AlternateContent":
            choice = child.find(f"{{{_MARKUP_COMPAT_NS}}}Choice")
            if choice is None:
                choice = child.find(f"{{{_MARKUP_COMPAT_NS}}}Fallback")
            if choice is not None:
                _pptx_walk_shapes(zf, slide_path, rels, choice, title, body)


def _pptx_slide_content(
    zf: zipfile.ZipFile,
    path: str,
    rels: dict[str, tuple[str, str, str]] | None = None,
) -> tuple[str, bool]:
    """A slide's text (title first, then shapes, tables, picture descriptions
    and chart values in shape-tree order) and whether the slide is hidden."""
    root = _read_zip_xml(zf, path)
    if root is None:
        return "", False
    if rels is None:
        rels = _relationship_entries(zf, _rels_path_for(path))
    p = f"{{{_PPTX_MAIN_NS}}}"
    tree = root.find(f"{p}cSld/{p}spTree")
    title: list[str] = []
    body: list[str] = []
    if tree is not None:
        _pptx_walk_shapes(zf, path, rels, tree, title, body)
    return "\n".join([*title, *body]), root.get("show") == "0"


def _pptx_slide_text(zf: zipfile.ZipFile, path: str) -> str:
    return _pptx_slide_content(zf, path)[0]


def _pptx_notes_text(
    zf: zipfile.ZipFile,
    path: str,
    rels: dict[str, tuple[str, str, str]] | None = None,
) -> str:
    """The speaker notes of a slide: the notes page's body placeholder only,
    not its slide image, slide number or footer placeholders."""
    if rels is None:
        rels = _relationship_entries(zf, _rels_path_for(path))
    p = f"{{{_PPTX_MAIN_NS}}}"
    for rel_type, target, mode in rels.values():
        if not rel_type.endswith("/notesSlide") or mode == "External":
            continue
        notes = _read_zip_xml(zf, _resolve_part_path(path, target))
        if notes is None:
            return ""
        lines: list[str] = []
        for shape in notes.iter(f"{p}sp"):
            placeholder = shape.find(f"{p}nvSpPr/{p}nvPr/{p}ph")
            if placeholder is not None and placeholder.get("type") == "body":
                lines.extend(_pptx_paragraph_lines(shape.find(f"{p}txBody")))
        return "\n".join(lines)
    return ""


def _parse_pptx(content: bytes, *, deadline: float | None = None, max_chars: int | None = None) -> str:
    """Slide text in presentation order, one section per slide, with speaker
    notes in a `[Notater]` block and hidden slides labelled. Past `deadline`
    (a `time.monotonic()` value), or once the text passes `max_chars`, the deck
    is refused."""
    token = _READ_DEADLINE.set(deadline)
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            _zip_budget(zf)
            sections = []
            total_chars = 0
            for index, path in enumerate(_pptx_slide_parts(zf), start=1):
                rels = _relationship_entries(zf, _rels_path_for(path))
                text, hidden = _pptx_slide_content(zf, path, rels)
                notes = _pptx_notes_text(zf, path, rels)
                lines = [f"[Slide {index}]" + (" (skjult)" if hidden else "")]
                if text:
                    lines.append(text)
                if notes:
                    lines.extend(["[Notater]", notes])
                sections.append("\n".join(lines))
                total_chars += len(sections[-1]) + 2
                if max_chars is not None and total_chars > max_chars:
                    raise OfficeArchiveTooLargeError("The document holds too much text")
            return "\n\n".join(sections).strip()
    except zipfile.BadZipFile as exc:
        raise _MalformedOfficeDocument("Not a valid .pptx archive") from exc
    finally:
        _READ_DEADLINE.reset(token)


def _parse_tika(
    filename: str,
    content: bytes,
    content_type: str | None,
    tika_url: str,
    timeout_s: float,
) -> str:
    headers = {
        "Accept": "text/plain",
        "Content-Type": content_type or "application/octet-stream",
        "Content-Disposition": f'attachment; filename="{Path(filename).name}"',
    }
    with httpx.Client(timeout=timeout_s) as client:
        response = client.put(
            f"{tika_url.rstrip('/')}/tika",
            content=content,
            headers=headers,
        )
        response.raise_for_status()
        return response.text.strip()


def _parse_local(
    extension: str,
    content: bytes,
    metadata: dict[str, Any],
    *,
    max_chars: int | None = None,
) -> str:
    if extension == "pdf":
        return _parse_pdf(content)
    if extension == "docx":
        return _parse_docx(content)
    if extension in {"txt", "md"}:
        return _parse_plain(content)
    if extension == "csv":
        return _parse_csv(content)
    if extension == "json":
        text, extra = _parse_json(content)
        metadata.update(extra)
        return text
    if extension in {"html", "htm"}:
        return _parse_html(content)
    if extension == "xlsx":
        # `max_chars` is a chat-turn limit; durable imports pass none and keep every row.
        return _parse_xlsx(content, max_chars=max_chars, info=metadata)
    if extension == "pptx":
        # Like a workbook, only a chat turn is latency-bound.
        deadline = time.monotonic() + _PARSE_DEADLINE_SECONDS if max_chars is not None else None
        return _parse_pptx(content, deadline=deadline, max_chars=_PPTX_MAX_CHAT_CHARS if max_chars is not None else None)
    raise UnsupportedFileTypeError(f"Unsupported file type: {extension}")


def parse_uploaded_file(
    filename: str,
    content: bytes,
    content_type: str | None,
) -> ImportDocument:
    extension = Path(filename).suffix.lower().lstrip(".")
    settings = get_settings()
    metadata: dict[str, Any] = {
        "filename": filename,
        "content_type": content_type,
        "size_bytes": len(content),
        "sha256": _sha256(content),
        "extension": extension,
    }

    text = ""
    last_tika_error: str | None = None
    processor_order = settings.document_processor_order or ["local"]
    for processor in processor_order:
        if processor == "tika":
            if not settings.tika_enabled:
                continue
            try:
                text = _parse_tika(
                    filename,
                    content,
                    content_type,
                    settings.tika_url,
                    settings.tika_timeout_s,
                )
                if text:
                    metadata["parser"] = "tika"
                    break
            except (httpx.HTTPError, OSError) as exc:
                last_tika_error = exc.__class__.__name__
                continue
        elif processor == "local":
            text = _parse_local(extension, content, metadata)
            metadata["parser"] = "local"
            break

    if not text and "local" not in processor_order:
        text = _parse_local(extension, content, metadata)
        metadata["parser"] = "local"

    if last_tika_error and metadata.get("parser") != "tika":
        metadata["tika_fallback_error"] = last_tika_error

    title = Path(filename).stem
    return ImportDocument(source_name=filename, title=title, text=text, metadata=metadata)
